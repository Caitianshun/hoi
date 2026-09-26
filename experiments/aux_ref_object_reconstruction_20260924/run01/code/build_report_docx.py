"""Portable concise report with embedded figures; run with bundled document Python."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
E=Path(__file__).resolve().parents[1]
data=json.loads((E/'report_data.json').read_text());summary=json.loads((E/'evaluation/cross_pose/summary.json').read_text())
D=Document();s=D.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=D.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');st.paragraph_format.space_after=Pt(6)
    grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(grid)
D.styles['Normal'].font.size=Pt(10.5);D.styles['Normal'].paragraph_format.line_spacing=1.13
D.styles['Title'].font.size=Pt(20);D.styles['Heading 1'].font.size=Pt(15);D.styles['Heading 2'].font.size=Pt(12);D.styles['Caption'].font.size=Pt(9)
D.styles['Caption'].font.bold=False;D.styles['Caption'].font.italic=False;D.styles['Subtitle'].font.italic=False
for st in D.styles:
    for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
def p(t,style=None):return D.add_paragraph(t,style)
def page(t):
    q=D.add_heading(t,level=1);q.paragraph_format.page_break_before=True
def table(rows,widths):
    t=D.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,(cell,txt,w) in enumerate(zip(cells,row,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for tag,value in [('top',75),('bottom',75),('left',85),('right',85)]:
                el=OxmlElement('w:'+tag);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');m.append(el)
            pr.append(m);b=OxmlElement('w:tcBorders')
            for tag in ['top','left','bottom','right']:
                el=OxmlElement('w:'+tag);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            pr.append(b)
            if i==0:
                el=OxmlElement('w:shd');el.set(qn('w:fill'),'E8EEF4');pr.append(el)
            q=cell.paragraphs[0];q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=1.05;q.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            r=q.add_run(str(txt));r.font.size=Pt(9.5);r.bold=i==0
        if i==0:
            el=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(el)
        el=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(el)
    gap=p('');gap.paragraph_format.space_after=Pt(3);gap.paragraph_format.line_spacing=Pt(3)
def image(path,width):
    q=p('');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(3)
    shape=q.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set('descr',path.stem.replace('_',' '))
D.core_properties.title='物体运动条件对高斯表示学习的辅助验证';D.core_properties.author='HOI research project'
p('物体运动条件对高斯表示学习的辅助验证','Title')
p('2026年9月26日  附件第14节  AUX_REF_OBJECT','Subtitle')
p(data['conclusion'])
p('本轮要区分两个变化：运动不准会直接把同一物体投到错误位置，也可能在训练中改变颜色、不透明度、尺度和表面附着。为分开观察这两者，每个事件重新学习两个物体表示，再分别放到两种固定运动下渲染。人体、背景和RGB人体运动始终取同一个旧S1并全部冻结。')
table([['协议','训练条件','物体运动来源','其余可训练参数'],['AUX_REF_OBJECT','Pred','旧S1最终RGB预测','仅新物体高斯属性'],['AUX_REF_OBJECT','Ref','发布fit01物体拟合','与Pred完全相同']],[1.45,.7,2.25,2.15])
p('Ref是额外信息条件，不是新算法，也不是无噪声真值；发布拟合可能利用多视角或RGB-D。P0主协议仍不读发布姿态。发布人体拟合、传感深度、纹理和camera1 RGB均不用于训练；发布人体fit02仅用于原有评价区域的栅格化。')
table([['协议','事件','训练S','评价E','正式步数'],['AUX_REF_OBJECT','箱体','12帧','5帧','Pred与Ref各8000'],['AUX_REF_OBJECT','木椅','9帧','4帧','Pred与Ref各8000']],[1.45,1.0,1.05,1.05,2.0])
p('S由原生camera0与同采集样本物体参数组成，E取其中已有camera1的采集组。没有把旧14/9个近邻评分槽变成训练标签；全部输入文件已核验身份。')
p('官方采集组身份确定，真实曝光及组内同步误差未知。Pred仅在端点内插值；Ref无插值或最近复制，无ICP或尺度对齐。坐标转换一致性核验通过，不等于拟合精度已获验证。','Caption')
p('共同新物体初值为4096个未训练高斯，旧物体外观未加载；初始颜色仍可能受旧Pred投影采色偏置。两臂同seed 12345、同帧顺序、同损失和6000点上限。只更新静态5mm范围附着、颜色、不透明度、尺度和方向；全部运动、ordinary track与运动先验关闭。')

page('统一保留视角的四格结果')
p('第一字母表示训练得到的高斯，第二字母表示渲染运动；P为Pred，R为Ref。例如PR是Pred训练出的表示放到Ref运动下。9帧全部保留，四终态先冻结，随后统一camera1评价；没有按成绩选终态、缩小掩码或挑帧。')
rows=[['协议','事件与格','PSNR dB','SSIM','LPIPS']]
for dev,name in [('dev1','箱体'),('dev2','木椅')]:
 for cell in ['PP','PR','RP','RR']:
  m=summary['per_dev'][dev]['cells'][cell]['object']
  rows.append(['AUX_REF_OBJECT',name+' '+cell,f'{m["psnr_db"]["mean"]:.3f}',f'{m["ssim"]["mean"]:.4f}',f'{m["lpips_spatial_mean"]["mean"]:.4f}'])
table(rows,[1.5,1.4,1.05,1.05,1.55])
p('表中为每帧固定物体区域指标的等权均值；PSNR和SSIM越高越好，LPIPS越低越好。LPIPS为现成本地AlexNet空间图在同一区域上的均值，非裁剪图标量分数。全部逐帧值、有效帧数、均值与中位数均保留在JSON/CSV。','Caption')
rows=[['协议','事件','共同Ref差均值','差的中位数','SSIM差均值']]
for dev,name in [('dev1','箱体'),('dev2','木椅')]:
 v=summary['per_dev'][dev]['differences']['learned_ref'];rows.append(['AUX_REF_OBJECT',name,f'{v["psnr_db"]["mean"]:+.3f} dB',f'{v["psnr_db"]["median"]:+.3f} dB',f'{v["ssim"]["mean"]:+.4f}'])
table(rows,[1.45,.7,1.65,1.6,1.15])
p('主差值RR−PR控制同一个Ref渲染运动，比较学到的表示。工程门槛要求两事件各自PSNR均值至少+0.5 dB、逐帧差中位数为正、SSIM均值差不低于−0.005，并通过透明化、核膨胀、越界和输入边界检查。这不是统计显著性，也不是方法晋级。')
p(data['cross_interpretation'])
p('逐格均值与逐帧差的中位数是不同统计量，不能用两格中位数相减替代后者。代数分解不代表独立因果贡献，本文不报告所谓污染贡献百分比。','Caption')

for dev,name in [('dev1','箱体'),('dev2','木椅')]:
 page(name+'全部评价样本的对照')
 image(E/f'figures/{dev}_all_E_grid.png',7.05)
 p('协议 AUX_REF_OBJECT。每行依次为真实RGB、PP、PR、RP、RR和只渲染冻结人体背景。青线是同一个固定物体区域。所有E帧均展示；裁剪只由GT物体框加80像素边距决定，同一行六列完全相同。每格下方为固定O区域PSNR与SSIM；完整640×480图另存。','Caption')
 p(data['visual_notes'][dev])
 if dev=='dev1':p('t26的评价物体区域只有200像素，仍保留且单独报告像素数。参考拟合投影区域存在拟合及同步误差，分数不能直接证明真实接触或几何正确。','Caption')
 else:p('训练camera0与评价camera1分开报告。旧人体与背景已使用完整camera0历史，因此本轮不能称为少量原生帧从零重建整个场景。','Caption')

page('实际表示与监督可达性的检查')
image(E/'figures/attribute_distributions.png',6.4)
p('协议 AUX_REF_OBJECT。箱线显示实际终态高斯属性的中位数、四分位区间和1.5倍IQR须；为可读性隐藏图中极端散点，但完整极值与高斯数组均已导出。红虚线为5mm规范偏移硬边界。拓扑允许因训练条件不同而不同，不能把各臂稳定ID直接当成材料真值对应。','Caption')
rows=[['协议','事件与臂','点数','非零数据梯度步','最大尺度mm','最大偏移mm']]
for key,r in data['runs'].items():
 rows.append(['AUX_REF_OBJECT',key.replace('dev1','箱体').replace('dev2','木椅'),str(r['points']),str(r['gradient_steps']),f'{r["max_scale_mm"]:.2f}',f'{r["max_offset_mm"]:.3f}'])
table(rows,[1.35,1.15,.65,1.6,1.05,.75])
p(data['representation_interpretation'])


page('固定场景限制与研究决定')
p(data['fixed_scene_interpretation'])
rows=[['协议','事件与条件','有效深度样本','T低于0.1','模板未覆盖O']]
for key,r in data['transmittance'].items():
 rows.append(['AUX_REF_OBJECT',key.replace('dev1','箱体').replace('dev2','木椅'),str(r['valid_samples']),f'{100*r["low_T_fraction"]:.1f}%' if r['low_T_fraction'] is not None else '未测',f'{100*r["uncovered_O_fraction"]:.1f}%' if r['uncovered_O_fraction'] is not None else '未测'])
table(rows,[1.4,1.35,1.35,1.15,1.3])
p('透射率T表示物体模板第一表面深度之前尚未被冻结人体/背景遮住的光路比例。每帧最多512个确定性固定O像素，T低于0.1的分母为其中有模板深度的样本；模板未覆盖率用全部O像素单列。按高斯中心深度排序和原alpha截断复算，不是把H/S整体alpha当成物体前遮挡。模板深度与实际高斯支撑并不完全等价，CPU核数值检查不能消除该几何近似。','Caption')
p('每步RGB/实例项梯度和静态正则梯度分开记录。非零数据梯度只说明观测项能更新参数，不能单凭这个计数证明更新方向正确或同一材料表面得到支持。箱体t21和木椅t8的冻结SAM2物体标签为空；仍保留并按原分母计算，未补步。')
p('共同H/S-only图、规范位置、锚点与子代关系、颜色、opacity、尺度、方向、5mm偏移，以及每步固定O贡献均已导出。可信局部材料对应未测，未建立新匹配器。','Caption')
D.add_heading('本轮决定',level=2);p(data['decision_text'])
page('实测成本和复算记录')
p(data['cost_text'])
p('新增位姿求解0次；正式高斯训练4次，每次8000步，预检查临时两条件各8步且未带入正式训练；独立事件0条。全部运动冻结。没有S2、S3、注意力、接触、新骨干、追加种子或全三支重训。')
p('目录 experiments/aux_ref_object_reconstruction_20260924/run01 保存PROTOCOL、逐样本清单、frozen_aux、实际检查点、逐步梯度、四格逐帧指标、诊断、REPRODUCE和NEXT_DECISION。协议和结果每张表均标AUX_REF_OBJECT；模型和数据不随代码上传。')
D.add_heading('完整性和独立复算',level=2)
p('211项资产身份一致，16个阶段检查点仅含物体五类优化参数，人体背景与RGB运动冻结，日程及初值配对相同。历史8份核心源码和4份初始化及旧S1检查点未变。36份float渲染独立复算PSNR、SSIM和LPIPS，135条逐帧差值及所有汇总一致。')
p('验收曾因CPU与CUDA的float32 sigmoid双重舍入触发1e−7断言，失败记录保留；只将验收参考改为float64计算，原容差不变。训练、冻结输入与模型输出均未修改。正式四次训练与评价无失败或重跑。')
p('新增原生数据仅21份物体fit01参数和9张camera0图像，约1.74 MB；模板中心化及R/t转换复算最大顶点差约0.22微米。该微小差值证明文件转换一致，不能解释为发布拟合接近真实几何的精度。','Caption')
p('阶段目标9月27日前完成；10月7日前收敛核心问题、11月4日冻结结果、11月5—15日完整写作窗口不变。本轮输入较强且只验证两个已开发事件，不能外推为跨序列泛化、长遮挡恢复或运动接触创新。','Caption')
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
out=E/'output/AUX_REF_OBJECT_verification.docx';out.parent.mkdir(exist_ok=True);D.save(out);print(out)
