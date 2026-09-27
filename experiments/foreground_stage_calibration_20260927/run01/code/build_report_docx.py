"""Portable V4 report, authored with bundled python-docx from private results."""
import json,csv,math,hashlib,importlib.util,zipfile,statistics
from pathlib import Path
from PIL import Image,ImageDraw,ImageFont
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

RUN=Path(__file__).resolve().parents[1];ROOT=RUN.parents[2]
OLD=ROOT/'experiments/baseline_protocol_calibration_20260927/run01'
spec=importlib.util.spec_from_file_location('v3_document_helpers',OLD/'code/build_report_docx.py');v3=importlib.util.module_from_spec(spec);spec.loader.exec_module(v3)
read=v3.read;num=v3.numeric
def rows(path):
    with Path(path).open() as f:return list(csv.DictReader(f))

class Report(v3.Report):
    def __init__(self,content,costs):
        super().__init__(RUN,content,costs)
        self.doc.core_properties.title='前景训练阶段与保留视角颜色诊断'
        self.doc.core_properties.subject='V4 执行结果与失败边界'
        for run in self.doc.sections[0].footer.paragraphs[0].runs:
            if 'V3' in run.text:run.text=run.text.replace('V3 基线校准','V4 前景阶段校准')
    def image(self,path,caption,width=None,max_height=None):
        path=Path(path);width=width or self.width
        with Image.open(path) as im:
            if max_height:width=min(width,max_height*im.width/im.height)
        p=self.p('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(3)
        shape=p.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set('descr',caption)
        self.p(caption,'Caption')
    def sheet(self,name,records,columns=1,strip_header=False):
        cards=[];font=ImageFont.truetype(v3.DEJAVU,34)
        for record in records:
            path=v3.asset(record)
            with Image.open(path) as src:im=src.convert('RGB')
            if strip_header:im=im.crop((0,36,im.width,im.height))
            width=1400;im=im.resize((width,round(im.height*width/im.width)),Image.Resampling.LANCZOS)
            card=Image.new('RGB',(width,im.height+54),'white');card.paste(im,(0,54))
            label=record.get('label',record.get('frame_id',path.stem));ImageDraw.Draw(card).text((8,8),label,font=font,fill='black');cards.append(card)
        cellh=max(x.height for x in cards);cellw=cards[0].width;gap=20
        sheet=Image.new('RGB',(columns*cellw+(columns-1)*gap,math.ceil(len(cards)/columns)*(cellh+gap)-gap),'white')
        for i,c in enumerate(cards):sheet.paste(c,((i%columns)*(cellw+gap),(i//columns)*(cellh+gap)))
        path=self.figure_dir/(name+'.jpg');sheet.save(path,quality=90,subsampling=0,optimize=True);return path

def run():
    content=read(RUN/'report_content.json');cost=read(RUN/'costs.json');ev=read(RUN/'evaluation/summary.json');app=read(RUN/'diagnostics/appearance/summary.json');support=read(RUN/'initialization_support.json');figs=read(RUN/'figure_manifest.json');verify=read(RUN/'protocol/verification.json');decision=read(RUN/'protocol/final_decision.json');ex=read(RUN/'protocol/fixed_examples.json')
    r=Report(content,cost);r.p('前景训练阶段与保留视角颜色诊断','Title');r.p('V4  执行结果与后续研究决定','Subtitle')
    for p in content['summary']:r.p(p)
    table=[['分支','实际状态','路径名义轮','路径更新','新终态']]
    for a in cost['formal_attempts']:table.append([a['run'],a['status'],a.get('nominal_iterations',a.get('attempted_iterations')),a['optimizer_updates'],'有' if a.get('checkpoint') else '无'])
    r.table(table,[1.1,2.09,1.1,1.2,1.5],10)
    r.p('表中路径计数不含作废与重复轮数；全部实际尝试见第5页。原 H1 为全图 L1 基线，W_fine 仅调整 fine 阶段，W_all 调整两个阶段。空值表示没有可比较的终态，不表示零分。')
    for p in content['decision_summary']:r.p(p)
    r.p('证据入口：本报告的绝对分数、配对差和成本读取 evaluation/summary.json、costs.json；逐帧原值、完整数组身份与失败记录见反馈包。','Caption')

    r.page('协议与损失的实际含义')
    r.p('Backpack 使用同一组 268 个训练帧、16 个开发保留帧及发布相机，分辨率 1277×718。时间由原帧号归一化，不能解释为实测秒。00000 位于训练时间范围外，保留在主表；另报排除它后的 15 帧作为补充。已查看过的 16 帧用于开发筛选，不能作为独立泛化确认。')
    r.p('每张完整合成图先计算逐像素 RGB 绝对误差均值，再分别对前景和背景的像素求平均，两项各乘 0.5；最后在 batch 内按图像等权平均。只有一个非空区域时，它获得全部权重。发布训练 mask 的阈值为 128，原始预测颜色不新增裁剪。损失仍更新同一个场景的高斯位置、形状、透明度、颜色与允许的形变参数。')
    r.p('这会提高每个前景像素的影响，同时经投影与渲染梯度影响增密和剪枝；相同密度规则不保证相同点拓扑。本轮测量的是监督分配及其诱导优化的总效应，不能把差异全部解释为某一种重影或某个独立模块。SSIM 和 LPIPS 不参与训练，原有三项 fine 正则保持。')
    r.p('W_fine 分叉加载 coarse 模型和 RNG，再按官方入口重置 fine Adam、梯度累计和采样栈，保留 max_radii2D。首批训练 ID 和纯全图 L1 加正则的无更新探针保存在 branch_validation.json；旧日志缺批次 ID，且 l1 别名包含正则，因此历史首步总损失无法直接配对。')
    r.p('新增训练 mask 只用于本轮训练误差权重，没有修改标签、初值或相机。人体与静放背包并非始终属于同一个前景区域；本轮没有独立人体和背包指标，也不以图像分数证明接触、材料点或三维几何正确。')
    r.table([['预算项','实际','上限'],['预定正式分支',cost['formal_attempt_count'],cost['limits']['formal_attempts']],['名义优化轮',cost['formal_nominal_attempts'],cost['limits']['nominal_steps']],['临时优化轮',cost['temporary_nominal_steps'],cost['limits']['temporary_steps']],['GPU 任务小时',num(cost['gpu_task_hours']),num(cost['limits']['GPU_seconds']/3600)]],[2.79,2.1,2.1],10)
    r.p('所有 GPU 任务绑定物理 GPU1 RTX 3090。GPU 时间是串行任务进程墙钟，包含加载、导出及失败；不是 CUDA 核计时。历史 H1 coarse 和其他 V3 成本单列，不冒记为零。')

    r.page('输入掩码与初始表面支持')
    for p in content['support']:r.p(p)
    q=support['stats'];tab=[['来源点','数量','重投影中位 px','视差中位 度','源帧间隔中位']]
    for pool,region in [('raw','foreground'),('raw','background'),('selected','foreground'),('selected','background')]:
        v=q[pool][region];tab.append([pool+' '+region,v['points'],num(v['reprojection_px']['p50']),num(v['parallax_deg']['p50']),num(v['source_frame_gap']['p50'],1)])
    r.table(tab,[2.09,1.0,1.3,1.3,1.3],9)
    cloud=next(x for x in support['figures'] if 'foreground_source_time' in x['path']);r.image(cloud['path'],'前景初始点按两次来源帧号的平均值着色。显示原始世界坐标，未将移动主体统一到规范空间。离散时刻的分布不能当作真实运动轨迹或几何误差真值。',max_height=3.8)

    r.page('固定训练帧的掩码与来源观测')
    records=[x for x in support['figures'] if x.get('selection')=='uniform8'];r.image(r.sheet('fixed_training_support',records,2),'时间等间隔选定的全部八帧。每幅保留原 RGB、发布 mask 与实际接受的源关键点；没有把其他时刻的全部点投到当前帧冒充材料对应。',max_height=7.7)
    r.p('源支持圆盘半径 8 像素，仅表示附近存在被接受的输入观测，不等于当前渲染覆盖、可靠表面或跨时刻身份。两个额外面积极值帧也已检查，图像随反馈包保存。','Caption')

    r.page('训练过程与失败证据')
    for p in content['training']:r.p(p)
    path=RUN/'diagnostics/training_summary.jpg'
    if path.exists():r.image(path,'每 100 轮记录的真实损失与点数，重叠恢复段另线绘制。批次随迭代改变，不能当作固定验证集性能。作废恢复隔离在失败记录中；未记录的逐步损失不插值。',max_height=4.2)
    r.p('W_all 使用附加的只读观察入口保存滚动检查点和非有限状态；冻结的训练函数、区域损失、学习率、随机采样、增密与优化器规则均未改。该修复保证后续异常可追踪，但不能补回已经退出的 W_fine 状态。')

    r.page('Backpack 全部帧的数值评价')
    tab=[['区域与集合','H1 PSNR','W fine PSNR','W all PSNR','W all 减 H1']]
    for split,sl in [('train','训练'),('retained','开发16'),('retained_without_00000','补充15')]:
        for reg,label in [('foreground','前景'),('background','背景'),('full','全图')]:
            vs=[ev['summaries'][b][split][reg]['psnr_db'] for b in ['H1','W_fine','W_all']];tab.append([sl+' '+label,*[num(v) for v in vs],num(None if vs[2] is None else vs[2]-vs[0],signed=True)])
    r.table(tab,[1.99,1.2,1.3,1.3,1.2],9)
    tab=[['分支相对 H1','前景均值 dB','前景中位 dB','前景 LPIPS 差','背景均值 dB']]
    for branch in ['W_fine','W_all']:
        d=ev['paired'][branch+'-H1']['retained'];tab.append([branch,num(d['foreground']['psnr_db']['mean'],signed=True),num(d['foreground']['psnr_db']['median'],signed=True),num(d['foreground']['lpips_spatial_mean']['mean'],signed=True),num(d['background']['psnr_db']['mean'],signed=True)])
    r.table(tab,[1.49,1.4,1.4,1.4,1.3],9)
    r.p('主门槛为开发16帧前景 PSNR 均值至少 +1 dB、逐帧差中位数大于 0、前景 LPIPS 至少下降 0.05、背景 PSNR 下降不超过 0.3 dB，并通过全部固定图例的质量检查。训练指标只用于解释拟合，不能替代开发门槛。')
    for p in content['hos']:r.p(p)
    r.p('SSIM 使用原 7×7 口径；LPIPS 使用固定 AlexNet 0.1 空间图，区域边界仍受感受野影响。SSE 为每像素先平均 RGB 通道平方误差再求和，MSE=SSE/像素数；CSV 同时保存逐帧均值与池化 PSNR 所需分母。显示与原主指标均裁剪到 [0,1]，raw 数值另外保留。','Caption')

    r.page('有限梯度探针与阶段解释')
    for p in content['probes']:r.p(p)
    pm=read(RUN/'diagnostics/state_probes/manifest.json');tab=[['快照','点数','opacity 中位','最大轴尺度 p95','位移 p95 中位']]
    for s in pm['snapshots']:
        vals=[x['world_displacement']['p95'] for x in s['deformation']];mid=statistics.median(vals) if vals else None
        tab.append([s['branch']+' '+s['stage']+' '+str(s['iteration']),s['attributes']['points'],num(s['attributes']['opacity']['p50']),num(s['attributes']['scale_max_axis']['p95']),num(mid)])
    r.table(tab,[2.69,1.0,1.1,1.1,1.1],8.6)
    r.p('四个训练探针帧固定为 '+', '.join(pm['probe_frame_ids'])+'。尺度和位移采用继承的 HOS 世界坐标单位，未核实其米制尺度。各参数组分别报告 FG/BG 梯度范数与夹角，不能跨组按范数排名。共同 coarse 起点只计算一次；未保存的快照标 NA。')
    path=RUN/'diagnostics/probe_summary.jpg'
    if path.exists():r.image(path,'固定四帧中，同一参数组的 FG 与 BG 梯度方向。未连接、零范数或非有限情况标 NA，不能由四帧梯度推导完整训练的因果机制。',max_height=3.8)

    r.page('BEHAVE 颜色与覆盖的冻结诊断')
    for p in content['appearance']:r.p(p)
    tab=[['序列与视图','native','只保留 DC','沿 camera0 求色']]
    for dev,label in [('dev1','箱体'),('dev2','木椅')]:
        for group,gname in [('camera1_E','保留 camera1'),('camera0_paired_E','配对 camera0')]:tab.append([label+' '+gname,*[num(app['summary'][dev][group][m]['full']['psnr_db']) for m in ['native','dc_only','cam0_direction']]])
    r.table(tab,[2.59,1.3,1.5,1.6],10)
    tab=[['序列与模式','RGB 大于1 通道%','任意越界像素%','低 alpha 像素%']]
    stats=read(RUN/'protocol/appearance_aggregates.json')
    for row in stats['rows']:tab.append([row['dev']+' '+row['mode'],num(100*row['raw_above1_channel_fraction'],2),num(100*row['raw_any_out_of_range_pixel_fraction'],2),num(100*row['alpha_below005_fraction'],2)])
    r.table(tab,[2.19,1.7,1.6,1.5],9)
    r.p('三模式只改变冻结模型颜色的求值。DC 是与高阶球谐 SH 系数共同学习的常数项，不等于从头训练 SH0。camera0 求色只替换 SH 所用中心，camera1 的投影、几何、排序和不透明度均保持。Alpha 通过白色常量点、黑背景渲染获得，是全场景累计不透明度，不是前景概率或几何真值。')

    # All predetermined figures remain visible even when a failed model is NA.
    for start in range(0,len(ex['hos_retained_frame_ids']),4):
        ids=ex['hos_retained_frame_ids'][start:start+4]
        r.page('开发保留全图 '+str(start//4+1));r.p('从左到右为 GT、H1、W_fine、W_all；灰色 NA 表示分支缺少终态。所有图同一尺寸，保留完整画面。','Caption')
        items=[x for fid in ids for x in figs['hos'] if x['frame_id']==fid and x['kind']=='full'];r.image(r.sheet('hos_full_'+str(start),items,1,True),'固定全部16帧的连续子集 '+', '.join(ids)+'。00000 是时间外推，仍在主表中。',max_height=7.8)
        r.page('开发前景同框裁剪 '+str(start//4+1));r.p('每组顺序为 GT、H1、W_fine、W_all。裁剪由训练前冻结的 mask 包围框加边距确定，各方法使用同一像素窗口。','Caption')
        items=[x for fid in ids for x in figs['hos'] if x['frame_id']==fid and x['kind']=='foreground_crop'];r.image(r.sheet('hos_crop_'+str(start),items,2,True),'同框裁剪 '+', '.join(ids)+'；分数仍来自完整的区域 mask，不是此矩形裁剪。',max_height=7.6)
    full=[x for x in figs['appearance'] if x['group']=='camera1_E' and x['kind']=='full comparison'];heat=[x for x in figs['appearance'] if x['group']=='camera1_E' and x['kind']=='alpha and overflow']
    for start in range(0,len(full),3):
        r.page('BEHAVE 固定保留视图 '+str(start//3+1));r.p('每行从左到右为 GT、native、只保留 DC、沿 camera0 方向求色。全部九个 E 按既定顺序展示。','Caption');r.image(r.sheet('behave_full_'+str(start),[dict(x,label=x['dev']+' '+x['frame_id']) for x in full[start:start+3]],1,True),'原始颜色统一裁剪到 [0,1] 用于展示；仅前向诊断，不作为新训练方法成绩。',max_height=7.8)
        r.page('BEHAVE 不透明度与颜色越界 '+str(start//3+1));r.p('从左到右为累计 alpha、native 超范围热图、DC 超范围热图、camera0 求色超范围热图。Alpha 灰度 0 到 1；热图固定显示 max(RGB)−1 的正部，色标 0 到 1，超过 1 饱和。','Caption');r.image(r.sheet('behave_alpha_'+str(start),[dict(x,label=x['dev']+' '+x['frame_id']) for x in heat[start:start+3]],1,True),'所有模式的 alpha、深度和可见性一致。真实超范围数值与通道/像素分母均在 CSV，热图饱和不改变原始数组。',max_height=7.6)
    r.page('同捕获组训练视图的诊断')
    camera0=[x for x in figs['appearance'] if x['group']=='camera0_paired_E'];r.image(r.sheet('camera0_selected',[dict(x,label=x['dev']+' '+x['frame_id']) for x in camera0],2,True),'每条序列按固定首、中、末 E 展示。每组为 GT、native、DC、camera0 求色。全部配对 camera0 的指标均在反馈包；不同相机区域标签来源和精确曝光同步限制保留。',max_height=7.6)

    r.page('下一步边界与复现入口')
    for p in content['next']:r.p(p)
    r.p('Tennis 仅完成 ZIP 中央目录与文件身份盘点。数值图像及 mask 一一匹配，未来划分按官方完整阶段的排序步长规则预声明；相机键与图像尺寸仍须实际验收。本轮未解码图像、产生模型分数或开始训练，不能以目录存在宣称完整协议已经通过。')
    r.p('本轮公开同步范围仅自产代码与源配置。结果 JSON、CSV、研究日志、原始附件、图像、模型及第三方源码副本保持本地。反馈包的固定小块保留 float32 原精度，包含 GT、mask、raw 预测及 alpha；缺失分支不补造。完整数组仍在训练机，由 SHA256 索引定位。')
    r.p('复算命令与安全前置条件见 REPRODUCE.md；不要直接重跑训练入口。核验器检查预算、模型与输入身份、全部配对键，并从保存的浮点图重新计算保留帧 PSNR/SSE。')
    r.p('相关一手来源','Heading 2')
    for source in read(RUN/'protocol/source_audit.json')['sources']:r.p(source['url']+'\n'+source['fact'],'Caption')
    r.p('锁定 Wu 4DGS 843d5ac636c37e4b611242287754f3d4ed150144；HOSNeRF ca169704a9d965e8266c034cceb506ffba72f405。HOSNeRF sample_subject_ratio=0.8 为 patch 采样比例，不能等同本轮完整图的 50/50 区域均值。准确文件哈希和后端身份见 run_manifest.json。','Caption')
    r.output.mkdir(exist_ok=True);path=r.output/'V4_foreground_stage_calibration.docx';r.doc.save(path)
    with zipfile.ZipFile(path) as z:
        media=[x for x in z.namelist() if x.startswith('word/media/')];external=[x for x in z.namelist() if x.endswith('.rels') and b'TargetMode="External"' in z.read(x)];assert media and not external
    manifest=dict(document=v3.ident(path),embedded_media=len(media),external_relationship_files=external,section_titles=r.section_pages,data_sources=[v3.ident(RUN/x) for x in ['report_content.json','costs.json','evaluation/summary.json','figure_manifest.json','protocol/verification.json','protocol/final_decision.json']],visual_QA='pending canonical render and every-page inspection')
    (RUN/'protocol/document_build.json').write_text(json.dumps(manifest,indent=2)+'\n');print(json.dumps(manifest['document']))
if __name__=='__main__':run()
