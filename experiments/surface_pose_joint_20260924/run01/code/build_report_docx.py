"""Concise editable report; run with the bundled document Python runtime."""
from pathlib import Path
import json,csv,datetime
import numpy as np
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
E=Path(__file__).resolve().parents[1];O=E/'output';O.mkdir(exist_ok=True);F=O/'figures';F.mkdir(exist_ok=True)
def load(p):return json.loads(Path(p).read_text())
M={d:{v:load(E/d/v/'metrics.json') for v in ['F00','F01','F10','F11']} for d in ['dev1','dev2']}
rows=list(csv.DictReader((E/'evaluation/paired_metrics.csv').open()));S={(x['dev'],x['method']):x for x in rows}
D=Document();section=D.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11);section.top_margin=section.bottom_margin=Inches(.65);section.left_margin=section.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=D.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
 st.paragraph_format.space_after=Pt(6)
 grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(grid)
D.styles['Normal'].font.size=Pt(11);D.styles['Normal'].paragraph_format.line_spacing=1.16
D.styles['Title'].font.size=Pt(20);D.styles['Heading 1'].font.size=Pt(15);D.styles['Heading 2'].font.size=Pt(12);D.styles['Caption'].font.size=Pt(9)
D.styles['Caption'].font.bold=False;D.styles['Caption'].font.italic=False;D.styles['Subtitle'].font.italic=False
for st in D.styles:
 for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
def p(t,style=None):return D.add_paragraph(t,style)
def head(t):return D.add_heading(t,level=1)
def page(t):
 q=head(t);q.paragraph_format.page_break_before=True
def table(data,widths=None):
 t=D.add_table(rows=0,cols=len(data[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,row in enumerate(data):
  cells=t.add_row().cells
  for j,(cell,txt) in enumerate(zip(cells,row)):
   if widths:cell.width=Inches(widths[j])
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for tag,value in [('top',70),('bottom',70),('left',80),('right',80)]:
    el=OxmlElement('w:'+tag);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');margins.append(el)
   pr.append(margins);b=OxmlElement('w:tcBorders')
   for tag in ['top','left','bottom','right']:
    el=OxmlElement('w:'+tag);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
   pr.append(b)
   if i==0:
    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8EEF4');pr.append(fill)
   q=cell.paragraphs[0];q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=1.05;q.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
   r=q.add_run(str(txt));r.font.size=Pt(9.5);r.bold=i==0
  if i==0:
   h=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(h)
  cant=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(cant)
 p('')
 return t
def image(path,width=6.9):
 q=D.add_paragraph();q.paragraph_format.space_after=Pt(3);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
 shape=q.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set('descr',Path(path).stem.replace('_',' '))
def fmt(x):return f'{float(x):.3f}'
D.core_properties.title='模板表面点与物体位姿联合优化验证';D.core_properties.author='HOI research project';D.core_properties.subject='Two development events controlled factorial pose experiment'
p('模板表面点与物体位姿联合优化验证','Title');p('2026年9月24日  两条开发序列  八组配对求解','Subtitle')
p('结论：投影图像项对相邻运动误差有有限改善，但没有变体通过完整晋级门槛。本轮停止原型扩展，保留旧S1；新增8000步高斯训练0次，独立事件消耗0条。')
p('问题是：可靠二维轨迹被绑定到可能错误的模板位置后，允许附着点沿表面调整、再加入随当前投影变化的图像残差，能否改善物体三维运动？八组全部从各自旧初始化开始，完成同一300步预算，冻结后统一评价。')
for d,name in [('dev1','箱体'),('dev2','木椅')]:
 D.add_heading(name+'位姿结果',level=2)
 data=[['方法','中心 cm','深度 cm','旋转 °','位移 cm','相对旋转 °']]
 for v,label in [('old_initialization','旧初始化'),('historical_P1','历史P1'),('F00','F00'),('F01','F01'),('F10','F10'),('F11','F11')]:
  x=S[d,v];data.append([label,*[fmt(x[k+'_mean']) for k in ['centroid_error_cm','camera_depth_absolute_error_cm','raw_rotation_error_deg','centre_displacement_error_cm','relative_rotation_error_deg']]])
 table(data,[1.2,1.1,1.1,1.1,1.1,1.45])
p('所有数值越小越好；深度为绝对误差，位移是相邻有效参考时刻的中心位移向量误差。中心使用同模板全部顶点均值，含历史孤立点。箱体14/14参考槽、13/13区间；木椅9/10槽、8/9区间，t=2秒缺失保留。输入仅camera0 RGB、原估计及固定无纹理模板；参考拟合、传感深度和camera1只评价。20%留出仍为同视频观测边，不是独立测试。', 'Caption')

page('共同输入与可验证实现')
table([['变体','共享表面点 q','当前投影图像项','检验目的'],['F00','固定','关闭','本轮共同基线'],['F01','固定','开启','H1 图像项作用'],['F10','可调','关闭','H2 表面点自由度'],['F11','可调','开启','两者联合及增量']],[.7,1.45,1.45,3.45])
p('观测池：箱体128轨迹，1274优化边、276留出边；木椅112轨迹，1112/250边。source由输入纹理和掩码内部距离选定，再用旧位姿一次射线求交。每条轨迹共享q，不强接片段；候选附着不等于真实材料身份。')
p('表面约束：q由原三角面及重心坐标表示，只能在初始q周围5 cm表面路径邻域移动。固定边上采样点提供保守路径上界，大三角面保留面内活动空间；每步最多跨一个相邻面，越界截断，不跨断开表面、不移动邻域中心或模板。')
p('图像项：RGB生成冻结的局部归一化颜色与梯度场。在当前R/t与q预测的像素处双线性采样，比较另一真实source的固定测得描述子。因此投影改变会改变残差。source图像边排除，二维约束保留；无新网络、自由描述子或扫描纹理。')
p('共同目标：保留原轮廓、弱估计深度、真实时间加速度项；新的二维项替换旧条件规范项，不重复叠加。二维尺度3 px；图像尺度由优化非source边的冻结MAD统计确定，下限0.15、系数1。预测遮挡、出界不会删掉困难边；另有共同越界/正深度约束。')
p('统一优化：全批量Adam；旋转向量/平移学习率均0.002，重心参数0.005，余弦降至初值0.1倍，300步或600秒取最终输出。本轮全部300步正常退出，未触发墙钟上限。F00更换了观测池和求解器，不能当作历史P1重放。')
p('实现核验：新旧基础项最大分项差2.75×10⁻⁶（含float32）。六条纹理边的R/t/q自动梯度与有限差分最大差1.86×10⁻⁶，固定target像素负对照梯度为0。source、留出、开关、固定分母及九项表面检查均通过。非零梯度不证明方向正确。')

page('二维拟合改善不等于三维恢复')
image(F/'input_diagnostics.png')
p('虚线为共同初值。释放q的F10大幅降低留出重投影残差，但没有形成两事件一致的三维增益；图像项F01相对F00的留出重投影反而略升。这说明新自由度能解释像素误差，仍不足以确认正确深度或材料身份。','Caption')
table([['事件及臂','平均q移动 cm','最大q移动 cm','曾触界点比例','面切换总次'],*[ [d+' '+v,fmt(M[d][v]['q_movement_mean_m']*100),fmt(M[d][v]['q_movement_max_m']*100),f"{100*(np.load(E/d/v/'surface_points.npz')['boundary_hit_count']>0).mean():.2f}%",M[d][v]['total_face_switches']] for d in M for v in ['F10','F11']]],[1.3,1.4,1.4,1.4,1.55])
p('所有最终q都在固定模板表面，证书路径距离不超过5 cm；固定q两臂没有面切换。可调点累计路径可能往返并超过5 cm，表中“移动”指相对初始点的欧式距离。触界包括邻域、模板边界及单步邻面上限，少数触界不证明扩大搜索有益。')
p('优化边有效覆盖：箱体F00/F01/F10/F11为98.51/98.35/95.60/96.08%，木椅94.78/94.60/93.97/94.24%。所有变体对F00降幅均小于5个百分点；没有大规模鲁棒拒绝。三维失败不能归因于通过删除观测作弊，但也不能因此归因于唯一机制。')

page('几何长尾与固定失败片段')
image(F/'reference_curves.png')
p('全部原参考槽公开，无ICP、尺度拟合或逐帧身份重定义。箱体后段大朝向误差基本保留；局部5 cm搜索和当前图像场未修复全局朝向分支。木椅参考缺失不补造，不把轨迹插值当真实遮挡恢复。','Caption')
table([['木椅64至74帧','旧初始化','历史P1','F00','F01','F10','F11'],['位移误差 cm','11.086','40.441','36.964','34.702','36.079','34.950'],['相对旋转 °','35.402','35.229','49.674','50.737','50.707','50.468']],[1.45,.9,.9,.9,.9,.9,1.1])
p('这个片段在上轮已暴露，本轮预先固定保留，不能称独立验证。F01/F11相对F00缓解位移，但距旧初始化仍有明显退化。F01木椅深度90分位由32.438升到34.296 cm，超过该比较的1.622 cm容差，平均收益没有消除长尾风险。')
p('参考来自拟合代理，箱体名义时间匹配最大约104 ms、木椅37.10 ms；两开发事件和单次配对不支持统计显著性、精细接触或全局不可恢复的结论。完整逐槽、区间和原始图像叠图随实验目录保留。')

page('方法决定与资源记录')
p('H1有限成立：固定q时，F01相对F00的相邻位移误差在箱体/木椅降低0.973/1.118 cm，超过0.845/0.872 cm预声明阈值。这是当前起点和预算下的局部相对运动收益；中心改善未超过两事件门槛，留出像素误差及木椅深度长尾也未同步改善。')
p('H2尚无统一支持：F10相对F00在箱体中心略差、两事件位移均更差，木椅中心改善不足以抵消运动退化。F11对F00没有同一主指标在两事件均超过阈值；差中差的位移变化为−0.029/−1.281 cm，中心为−0.088/+0.216 cm，不能据此宣称稳定协同。')
p('拒绝晋级的决定性原因：F01相对旧初始化，箱体中心/深度/位移增加5.642/6.299/2.652 cm，木椅位移增加3.991 cm；F10/F11也未通过旧初始化护栏。三维门槛未满足，所以没有追加高斯重建，不存在本轮camera1、实际高斯、查询或接触改善结论。')
table([['阶段','实测资源与耗时'],['观测池构造','CPU 0.38秒，不含历史跟踪前端'],['冻结RGB特征构造','箱体2.85秒，木椅2.54秒；复用原RGB'],['一次20步预运行','RTX3090 1.31秒求解'],['八臂正式求解','各12.67至21.55秒；累计约137.7秒求解'],['加载及导出','八进程总墙钟163.1秒；并非首次数据成本'],['完整runner及评价','173.61秒；独立评价0.85秒'],['显存及预算','峰值allocated 0.875 GiB；全部300步；高斯训练0次']],[1.65,5.4])
p('保留旧S1统一基线及所有失败记录；不按事件拼最好结果，不扩大局部邻域、不另开注意力、接触或新网络。本轮于9月24日完成阶段判断，早于9月27日界限。10月7日前据累计证据收敛主问题，11月4日核心结果冻结、11月5至15日完整写作窗口不变。')
p('可复算目录：experiments/surface_pose_joint_20260924/run01。REPRODUCE.md映射实际命令；protocol/frozen_surface_pose.json固定参数；evaluation/paired_metrics.csv及promotion_decision.json保存数字与判据；dev1、dev2目录保留全部位姿、表面点和优化过程。历史文件哈希与最终产物完整性单独核验。','Caption')
path=O/'Surface_pose_joint_verification.docx';D.save(path);print(path)
