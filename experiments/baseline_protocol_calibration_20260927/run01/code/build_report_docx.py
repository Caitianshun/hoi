"""Build the portable Chinese V3 report from frozen result files only.

Run with the bundled document Python. The caller must execute the document
skill's mark_artifact_operation_started.mjs exactly once before first creation.
This script never runs models, optimizers, metrics, or reads new test RGB. It
embeds previously generated comparison PNGs. Rendering/visual QA is separate.

Required report_content.json arrays: summary_paragraphs, behave_interpretation,
hos_interpretation, next_decision, limitations. Optional: title, subtitle,
coverage_interpretation, incidents, sources=[{label,url,detail}].
Required costs.json with optional fields: totals=[{label,value,unit,note}],
preprocessing=[{name,seconds,device,note}], notes=[str]. Formal run costs always
come directly from terminal run.json, not from narrative numbers. Package the
finished and visually verified report with code/package_results.py separately.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import math
import re
import zipfile
from datetime import datetime,timezone
from pathlib import Path

from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor

RUN=Path(__file__).resolve().parents[1]
ROOT=RUN.parents[2]
FONT='Noto Sans CJK SC'
REGIONS=['full','human','object','background','foreground']
REGION_NAMES={'full':'完整图','human':'人体 H','object':'物体 O','background':'背景 S','foreground':'前景 H并O'}
DEV_NAMES={'dev1':'箱体','dev2':'木椅'}
METHOD_NAMES={'E0':'完整 S1','4DGS':'适配 4DGS'}
DEJAVU='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'


def read(path):
    return json.loads(Path(path).read_text())


def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        while b:=f.read(1<<20):h.update(b)
    return h.hexdigest()


def ident(path):
    p=Path(path).resolve()
    return dict(path=str(p),sha256=sha(p),bytes=p.stat().st_size)


def asset(value):
    p=Path(value['path'])
    assert p.is_file(),p
    assert sha(p)==value['sha256'],f'Changed frozen figure {p}'
    return p


def numeric(x,d=3,signed=False):
    if x is None or isinstance(x,str) and not x.strip():return 'NA'
    value=float(x)
    if not math.isfinite(value):return 'NA'
    return f'{value:+.{d}f}' if signed else f'{value:.{d}f}'


def csv_rows(path):
    with Path(path).open() as f:return list(csv.DictReader(f))


def half_width(values):
    # Full checksums stay copyable while naturally wrapping in Word paragraphs.
    return '  '.join(values)


def unique_index(rows,fields,label):
    keys=[tuple(row[field] for field in fields) for row in rows]
    assert len(keys)==len(set(keys)),f'Duplicate records in {label}'
    return set(keys)


def validate_coverage(run,B,H,fmanifest):
    """Check saved metadata/CSV coverage before any comparison PNG is opened.

    Counts alone can hide one duplicate plus one missing frame. Validate exact
    frozen IDs for all E, paired input, full training metrics and HOS test rows.
    This does not recompute metrics or decode source GT/render arrays.
    """
    assert B['status']=='completed' and H['status']=='completed'
    assert B['same_input_system_comparison_not_single_variable_ablation'] is True
    assert B['ref_conditioned_B_F_excluded'] is True
    assert H['direct_fair_H0_H1_delta'] is None
    figures=fmanifest['figures'];expected_figures=set();expected_metrics=set();expected_fit=set();e_ids={}
    for dev,n in [('dev1',114),('dev2',98)]:
        base=run/'inputs'/f'behave_{dev}'
        train=read(base/'manifest.json')['frames'];ev=read(base/'evaluation_manifest.json')
        assert len(train)==n and len(unique_index(train,['frame_id'],dev+' train'))==n
        expected_e=5 if dev=='dev1' else 4
        groups={'camera1_E':ev['frames'],'camera0_paired_E':ev['paired_camera0_frames'],'camera0_full_training_fit':train}
        assert len(ev['frames'])==len(ev['paired_camera0_frames'])==expected_e
        assert {r['frame_id'] for r in ev['frames']}=={r['frame_id'] for r in ev['paired_camera0_frames']}
        e_ids[dev]=[r['frame_id'] for r in ev['frames']]
        for group,frames in groups.items():
            unique_index(frames,['frame_id'],dev+' '+group)
            preview=frames
            if group=='camera0_full_training_fit':
                # Same uniform floor indices as both frozen renderer producers.
                preview=[frames[i*(n-1)//15] for i in range(16)]
            expected_figures.update((dev,group,r['frame_id']) for r in preview)
            target=expected_metrics if group=='camera1_E' else expected_fit
            target.update((dev,group,r['frame_id'],method,region) for r in frames for method in ['E0','4DGS'] for region in REGIONS)
            for method in ['E0','4DGS']:
                for region in REGIONS:
                    d=B['results'][dev][group]['methods'][method][region]
                    for metric in ['psnr_db','ssim','lpips_spatial_mean']:
                        assert d[metric]['requested_frames']==len(frames)
                        assert 0<=d[metric]['valid_frames']<=len(frames)
    assert unique_index(figures,['dev','group','frame_id'],'BEHAVE figures')==expected_figures
    for row in figures:
        assert row['methods']==['GT','E0','4DGS'] and row['displayed_clipped_to_unit_range'] is True
        if row['group']=='camera1_E':
            assert row.get('crop') and len(row['crop_bounds_xyxy'])==4,'Every E requires its fixed crop'
            x0,y0,x1,y1=row['crop_bounds_xyxy'];assert 0<=x0<x1<=640 and 0<=y0<y1<=480
    fields=['dev','group','frame_id','method','region']
    for name,expected in [('metrics_per_frame.csv',expected_metrics),('input_fit.csv',expected_fit)]:
        rows=csv_rows(run/'evaluation/comparison'/name)
        assert unique_index(rows,fields,'BEHAVE '+name)==expected
    assert B['metric_rows']==len(expected_metrics)==90 and B['input_fit_rows']==len(expected_fit)==2210
    hos_base=run/'inputs/hos_backpack';hos_test=read(hos_base/'evaluation_manifest.json')['frames'];hos_train=read(hos_base/'manifest.json')['frames']
    test_ids={r['frame_id'] for r in hos_test};train_ids={r['frame_id'] for r in hos_train}
    assert len(hos_test)==len(test_ids)==16 and len(hos_train)==len(train_ids)==268 and not test_ids&train_ids
    hfig_ids=[Path(r['path']).stem for r in H['figures']]
    assert len(hfig_ids)==len(set(hfig_ids))==H['figure_count']==16 and set(hfig_ids)==test_ids
    for filename,group,ids,methods in [('metrics_per_frame.csv','test',test_ids,['H0','H1']),('input_fit.csv','input_fit',train_ids,['H1'])]:
        rows=csv_rows(run/'evaluation/hos_comparison'/filename)
        expected={(group,fid,method,region) for fid in ids for method in methods for region in ['full','foreground','background']}
        assert unique_index(rows,['group','frame_id','method','region'],'HOS '+filename)==expected
        for method in methods:
            for region in ['full','foreground','background']:
                d=H['results'][method+'_'+group][region]
                valid=sum(int(r['pixels'])>0 and r['sse_rgb_mean']!='' for r in rows if r['method']==method and r['region']==region)
                assert d['frames']==len(ids) and d['valid_frames']==valid
    assert H['native_H0_metrics']['frames']==16
    assert H['time_boundary_diagnostic']['frame_id'] in test_ids
    return dict(behave_E_ids=e_ids,behave_E_full_count=9,behave_E_crop_count=9,
                behave_input_preview_count=32,behave_paired_E_figure_count=9,hos_test_ids=sorted(test_ids),hos_test_figure_count=16)


class Report:
    def __init__(self,run,content,costs):
        self.run=run;self.content=content;self.costs=costs
        self.output=run/'output';self.figure_dir=self.output/'report_figures'
        self.figure_dir.mkdir(parents=True,exist_ok=True)
        self.image_sources=[];self.source_files=[];self.section_pages=[]
        self.doc=Document();sec=self.doc.sections[0]
        sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
        sec.top_margin=sec.bottom_margin=Inches(.59)
        sec.left_margin=sec.right_margin=Inches(.64)
        sec.footer_distance=Inches(.22)
        self.width=6.99
        for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption','Footer']:
            st=self.doc.styles[name];st.font.name=FONT;st.font.color.rgb=RGBColor(0,0,0)
            rf=st.element.get_or_add_rPr().get_or_add_rFonts()
            for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),FONT)
            st.paragraph_format.space_after=Pt(6)
        self.doc.styles['Normal'].font.size=Pt(10)
        self.doc.styles['Normal'].paragraph_format.line_spacing=1.12
        self.doc.styles['Normal'].paragraph_format.widow_control=True
        self.doc.styles['Title'].font.size=Pt(20)
        self.doc.styles['Subtitle'].font.size=Pt(10.5)
        self.doc.styles['Subtitle'].font.italic=False
        self.doc.styles['Heading 1'].font.size=Pt(15)
        self.doc.styles['Heading 2'].font.size=Pt(11.5)
        self.doc.styles['Caption'].font.size=Pt(8.6)
        self.doc.styles['Caption'].font.italic=False
        self.doc.styles['Caption'].paragraph_format.line_spacing=1.03
        self.doc.styles['Footer'].font.size=Pt(8)
        for st in self.doc.styles:
            for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        foot.add_run('V3 基线校准  ')
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
        self.doc.core_properties.title=content.get('title','动态人 物 场景重建基线校准')
        self.doc.core_properties.subject='V3 系统对照与数据协议执行总结'
        self.doc.core_properties.author='HOI 研究项目'

    def p(self,text,style=None):
        assert not re.search(r'cite|turn\d+(view|search)',text),'Tool citation token must not enter DOCX'
        return self.doc.add_paragraph(text,style)

    def paras(self,key,required=False):
        values=self.content.get(key,[])
        if required:assert values and all(isinstance(x,str) and x.strip() for x in values),key
        for value in values:self.p(value)

    def page(self,title):
        assert not any(c in title for c in ':：/—|'),'Headings use words/numbers/spaces only'
        p=self.doc.add_heading(title,1);p.paragraph_format.page_break_before=True
        p.paragraph_format.space_before=Pt(0)
        self.section_pages.append(title)

    def table(self,rows,widths,size=8.6,left_columns=(0,)):
        assert abs(sum(widths)-self.width)<.03,(sum(widths),self.width)
        tab=self.doc.add_table(rows=0,cols=len(widths));tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
        for col,width in zip(tab.columns,widths):col.width=Inches(width)
        for ri,row in enumerate(rows):
            assert len(row)==len(widths)
            cells=tab.add_row().cells
            for ci,(cell,value,width) in enumerate(zip(cells,row,widths)):
                cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                for side,val in [('top',60),('bottom',60),('left',75),('right',75)]:
                    el=OxmlElement('w:'+side);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
                pr.append(margins);borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                pr.append(borders)
                if ri==0:
                    el=OxmlElement('w:shd');el.set(qn('w:fill'),'E8EEF4');pr.append(el)
                paragraph=cell.paragraphs[0];paragraph.paragraph_format.space_after=Pt(0)
                paragraph.paragraph_format.line_spacing=1.04
                paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci in left_columns and ri else WD_ALIGN_PARAGRAPH.CENTER
                run=paragraph.add_run(str(value));run.font.size=Pt(size);run.bold=(ri==0)
            props=tab.rows[-1]._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
            if ri==0:props.append(OxmlElement('w:tblHeader'))
        after=self.p('');after.paragraph_format.space_after=Pt(0);after.paragraph_format.space_before=Pt(0)
        after.paragraph_format.line_spacing=Pt(3);after.paragraph_format.space_after=Pt(3)
        return tab

    def image(self,path,caption,width=None,max_height=None):
        pth=Path(path);width=width or self.width
        with Image.open(pth) as img:
            if max_height is not None:width=min(width,max_height*img.width/img.height)
        p=self.p('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.keep_with_next=True
        inline=p.add_run().add_picture(str(pth),width=Inches(width));inline._inline.docPr.set('descr',caption)
        cp=self.p(caption,'Caption');cp.paragraph_format.keep_together=True

    def contact(self,name,records,columns=1,mode='behave_full'):
        """Use only saved figure RGB; strip embedded labels then relabel legibly.

        No source image content is cropped in full mode: only prior figure label
        bands are removed. Crop mode retains the original fixed O crop image.
        """
        cards=[];font=ImageFont.truetype(DEJAVU,52)
        for record in records:
            ref=record['crop'] if mode=='behave_crop' else record.get('full',record)
            path=asset(ref)
            with Image.open(path) as im:im=im.convert('RGB')
            if mode=='behave_full':
                assert im.size==(1920,560),f'Unexpected BEHAVE figure geometry: {path}'
                content=im.crop((0,36,im.width,im.height-44))
            elif mode=='behave_crop':
                assert im.size==(960,300),f'Unexpected BEHAVE crop geometry: {path}'
                content=im.crop((0,36,im.width,im.height))
            elif mode=='hos_full':
                assert im.size==(1920,round(718*640/1277)+66),f'Unexpected HOS figure geometry: {path}'
                content=im.crop((0,66,im.width,im.height))
            else:raise ValueError(mode)
            label=record.get('report_label',record.get('frame_id',path.stem))
            target_width=1920 if mode!='behave_crop' else 960
            if content.width!=target_width:content=content.resize((target_width,round(content.height*target_width/content.width)),Image.Resampling.LANCZOS)
            label_height=76 if mode!='behave_crop' else 62
            card=Image.new('RGB',(target_width,content.height+label_height),'white')
            card.paste(content,(0,label_height));draw=ImageDraw.Draw(card)
            usefont=font if mode!='behave_crop' else ImageFont.truetype(DEJAVU,38)
            draw.text((16,6),label,font=usefont,fill='black')
            cards.append(card)
            self.image_sources.append(dict(report_contact=name,source_figure=ident(path),mode=mode,
                dev=record.get('dev'),group=record.get('group'),frame_id=record.get('frame_id',path.stem)))
        assert cards
        cellw=max(c.width for c in cards);cellh=max(c.height for c in cards)
        padding=24;rows=math.ceil(len(cards)/columns)
        sheet=Image.new('RGB',(columns*cellw+(columns-1)*padding,rows*cellh+(rows-1)*padding),'white')
        for i,card in enumerate(cards):sheet.paste(card,((i%columns)*(cellw+padding),(i//columns)*(cellh+padding)))
        path=self.figure_dir/f'{name}.png';sheet.save(path)
        return path

    def summary(self,B,A):
        self.p(self.content.get('title','动态人 物 场景重建基线校准'),'Title')
        self.p(self.content.get('subtitle','V3 执行总结  两条 BEHAVE 开发事件与 HOSNeRF Backpack 技术先导'),'Subtitle')
        self.paras('summary_paragraphs',True)
        self.p('本轮要区分三个原因：模型在输入图像上未充分拟合、输入视角没有覆盖保留视角的真实表面，以及人和物的运动或组合建模不足。先恢复完整旧 S1，再运行具有正常优化日程的公共骨干，能够缩小问题范围；系统差异本身不证明某个模块的因果贡献。')
        rows=[['事件','系统','全图均值 dB','全图 pooled dB','物体均值 dB','物体 pooled dB']]
        for dev in ['dev1','dev2']:
            for method in ['E0','4DGS']:
                rs=B['results'][dev]['camera1_E']['methods'][method]
                rows.append([DEV_NAMES[dev],METHOD_NAMES[method],numeric(rs['full']['psnr_db']['mean']),numeric(rs['full']['pooled_psnr_db']),numeric(rs['object']['psnr_db']['mean']),numeric(rs['object']['pooled_psnr_db'])])
        self.table(rows,[.7,1.1,1.3,1.3,1.3,1.29])
        self.p('主表仅包含完整 S1 原 RGB 预测运动与官方 Wu 4DGS 适配版。两者共享原 camera0 输入及同源合法初始先验，使用固定 camera1 E；S1 持续施加结构化先验，4DGS 自由形变和增长，因此是系统校准，不是单变量消融。','Caption')
        self.paras('behave_interpretation',True)
        self.p('PSNR 由图像均方误差换算，SSIM 反映局部结构，LPIPS 衡量感知差异；PSNR 和 SSIM 越高越好，LPIPS 越低越好。图像指标不能单独证明接触、相对位姿或物理几何正确。全部九帧及负结果保留，九帧不当作九个独立场景。')

    def protocol(self):
        self.page('输入协议与实现验收')
        rows=[['事件','训练帧','E 帧','训练时间秒','初始 H O S 点数','extent 米']]
        for dev in ['dev1','dev2']:
            m=read(self.run/'inputs'/f'behave_{dev}'/'manifest.json');s=read(self.run/'inputs'/f'behave_{dev}'/'shared_init.json')
            t=m['time_normalization'];counts=s['counts']
            rows.append([DEV_NAMES[dev],len(m['frames']),5 if dev=='dev1' else 4,f"{t['start_seconds']:.6f}\n{t['end_seconds']:.6f}",f"{counts['human']} / {counts['object']}\n{counts['background']}",numeric(s['scene_extent'],6)])
        self.table(rows,[.65,.7,.6,1.55,2.0,1.49])
        self.p('BEHAVE 使用原 640×480 去畸变 RGB、完整相机内参 K 和米制统一世界坐标。训练清单只列 camera0。归一化时间由训练起止时刻决定，保留相机同一名义时刻使用相同编码；完整发布人体与物体拟合、传感深度和 camera1 RGB 不进入新训练。')
        self.p('共享初始化直接来自旧 S1 未训练状态，保留人体、刚性物体和背景全部点及合法初始颜色。人体来自通用 SMPL X 与 RGB 预测姿态，物体使用已允许的无纹理几何和 RGB 初始运动，背景来自输入预测深度。没有从训练后 S1 或 AUX 取点或颜色。固定相机中心半径为零，因此 extent 采用点云逐轴中位中心到各点距离第 95 百分位的 1.1 倍；点、相机与形变网格未单独缩放。')
        rows=[['事件','原生渲染 RGB 最大差','原生 alpha 最大差','投影最大误差 像素']]
        for dev in ['dev1','dev2']:
            n=read(self.run/'protocol'/f'behave_{dev}_native_render_check.json')
            pr=read(self.run/'runs'/f'behave_{dev}_check'/'projection_check.json')
            rows.append([DEV_NAMES[dev],max(r['render_max_abs']['rgb'] for r in n['checks']),max(r['render_max_abs']['alpha'] for r in n['checks']),numeric(pr['max_pixel_error'],8)])
        self.table(rows,[.7,2.1,2.1,2.09])
        self.p('完整 S1 加载全部旧 H O S 状态和旧物体 bank。原生首 中 末时刻与原渲染同后端核对；中间时刻对物体和人体关节旋转做 SO 3 球面插值，平移及人体有界局部残差做线性插值，不使用参考运动补帧，超出原时间范围直接拒绝。','Caption')
        h=read(self.run/'inputs/hos_backpack/manifest.json');sampling=read(self.run/'inputs/hos_backpack/sampling_and_support.json')
        self.p(f"Backpack 使用完整阶段 loader 的实际划分，共 {h['count']} 个训练帧及 16 个测试帧，保留原生 {h['resolution'][0]}×{h['resolution'][1]} 尺寸。H1 初始化匹配、三角化与采色只读取冻结训练 RGB 和训练侧掩码；固定种子分层保留 {sampling['selected_points']:,} 点，其中背景 {sampling['selected_background']:,}、前景 {sampling['selected_foreground']:,}。发布相机可能由完整视频估计，这条信息路径单独披露，不等于点云颜色来自完整视频。")
        self.p('H0 使用官方完整阶段检查点及其原生人体 相机 状态预处理。检查点没有嵌入足以独立重建历史训练划分的完整身份；H0 作为渲染正控和原生条件参照。H0 与 H1 即使测试图像 ID 对齐，也不作公平的直接优劣差值。')

    def old_errors(self,A):
        self.page('历史结果的整图误差归因')
        self.p('先对历史 45 个 V2 浮点输出做零训练复算。H 人体、O 物体和 S 背景互斥且覆盖整图；先把每个像素的 RGB 均方误差在区域内求和得到 SSE，再求全图误差份额。dB 不能直接按区域面积加权。')
        rows=[['事件版本','全图均值','全图 pooled','H 份额 %','O 份额 %','S 份额 %','O 归零增量']]
        for dev in ['dev1','dev2']:
            for variant in ['B0','B1','F0','F1','F2']:
                d=A['camera1_existing'][f'{dev}_{variant}']
                rows.append([DEV_NAMES[dev]+' '+variant,numeric(d['full']['mean_frame_psnr_db']),numeric(d['full']['pooled_psnr_db']),numeric(100*d['H']['pooled_error_share'],2),numeric(100*d['O']['pooled_error_share'],2),numeric(100*d['S']['pooled_error_share'],2),numeric(d['algebraic_O_zero']['pooled_full_gain_db'],4)])
        self.table(rows,[1.05,1.0,1.05,.85,.85,.85,1.34])
        self.p('PSNR 和增量单位均为 dB；误差份额由整事件累加 SSE 后求比值。最后一列只把固定 O 内误差代数设零、保留 H S 误差，不是新的模型结果，也不是物体改进的物理上界，因为物体更新可能影响 O 外像素。','Caption')
        self.paras('coverage_interpretation')
        self.p(f"本次 {A['verification']['compared_historical_PSNR_values']} 项历史 PSNR 核对的最大绝对差为 {A['verification']['maximum_absolute_PSNR_difference_db']:.3g} dB。该一致性说明统计来源和原浮点结果吻合，不能独立解释误差的成因。背景误差份额大可能涉及覆盖、表示或优化，需要结合完整 S1 与公共基线的输入拟合来判断。")
        self.p('B0 B1 F 系列含 Ref 物体运动，最近物体监督只有箱体 12 和木椅 9 个原生时刻，人体与背景仍来自用过全部 114 和 98 帧的 S1。不能把整个系统描述成只见过 12 或 9 帧，也不能将这些强 Ref 辅助分支混入本轮同输入主表。')
        self.p('V2 的停止结论保留：独立可见监督没有达到两事件各增加 0.5 dB 的门槛，F2 相对同信息 F1 没有达到两事件各增加 0.2 dB 并有正中位差的条件。本轮系统校准不追改旧门槛，不把适配或初始化修复作为注意力必要性的证据。')

    def main_regions(self,B):
        self.page('保留视角的完整分区结果')
        rows=[['事件','系统','区域','均值 PSNR','pooled PSNR','SSIM','LPIPS']]
        for dev in ['dev1','dev2']:
            for method in ['E0','4DGS']:
                for region in REGIONS:
                    d=B['results'][dev]['camera1_E']['methods'][method][region]
                    rows.append([DEV_NAMES[dev],METHOD_NAMES[method],REGION_NAMES[region],numeric(d['psnr_db']['mean']),numeric(d['pooled_psnr_db']),numeric(d['ssim']['mean'],4),numeric(d['lpips_spatial_mean']['mean'],4)])
        self.table(rows,[.55,1.08,.95,1.15,1.25,1.0,1.01],8.2)
        self.p('每帧先算 PSNR 再等权平均，区别于先累加区域误差和像素再得到 pooled PSNR；两者同时给出。SSIM 为整图 7×7 映射在固定区域内的均值；LPIPS 为本地 AlexNet v0.1 空间映射区域均值，其感受野可跨区域边界。空区域记 NA。浮点预测原样保存，统一指标和展示均裁到 0 至 1。','Caption')
        if not B['lpips']['available']:
            self.p('本次统一 LPIPS 未取得有效值，表中记 NA，不解释为零误差。评价程序记录的原因为 '+B['lpips'].get('reason','未提供具体原因')+'。','Caption')
        self.p('camera1 仍使用历史发布拟合派生 H O S 区域；它不是独立人工 RGB 轮廓或真实可见性标注。箱体 t26 的 200 像素 O、困难帧及全部失败均保留。前景指 H 与 O 并集，不要求没有实体 bank 的 4DGS 产生虚构的实例贡献率。')

    def fit_and_pairs(self,B):
        self.page('输入拟合与逐帧配对变化')
        rows=[['事件','系统与时间集合','全图均值','全图 pooled','物体均值','物体 pooled']]
        for dev in ['dev1','dev2']:
            for group,label in [('camera0_full_training_fit','原训练全集'),('camera0_paired_E','E 名义时刻')]:
                for method in ['E0','4DGS']:
                    d=B['results'][dev][group]['methods'][method]
                    rows.append([DEV_NAMES[dev],METHOD_NAMES[method]+'\n'+label,numeric(d['full']['psnr_db']['mean']),numeric(d['full']['pooled_psnr_db']),numeric(d['object']['psnr_db']['mean']),numeric(d['object']['pooled_psnr_db'])])
        self.table(rows,[.55,1.84,1.15,1.15,1.15,1.15],8.4)
        self.p('原训练全集是训练集重建，不是泛化。配对 E 采用官方相同捕获组的真实 camera0 RGB，重新核对去畸变后逐像素一致；它们不是旧 S1 原训练帧。两者的时间来源不同，精确曝光同步未核实；camera0 区域来自 SAM2、camera1 区域来自拟合，均值差不能解释为纯视角损失。','Caption')
        rows=[['事件','区域','配对 PSNR 均值差','中位差','SSIM 均值差','LPIPS 均值差']]
        for dev in ['dev1','dev2']:
            for region in ['full','human','object','background','foreground']:
                d=B['results'][dev]['camera1_E']['paired_4DGS_minus_E0'][region]
                rows.append([DEV_NAMES[dev],REGION_NAMES[region],numeric(d['psnr_db']['mean'],3,True),numeric(d['psnr_db']['median'],3,True),numeric(d['ssim']['mean'],4,True),numeric(d['lpips_spatial_mean']['mean'],4,True)])
        self.table(rows,[.55,1.0,1.55,1.05,1.45,1.39],8.2)
        self.p('差值全部为适配 4DGS 减完整 S1，先逐帧配对再汇总。LPIPS 的负差才表示改善；中位数用于防止仅由少数大增益帧推动均值。两事件的结果不能证明未见序列泛化或全部 HOI 场景有效。','Caption')

    def behave_figures(self,figures):
        for dev in ['dev1','dev2']:
            self.page(DEV_NAMES[dev]+'全部保留视角图像')
            records=sorted([r for r in figures if r['dev']==dev and r['group']=='camera1_E'],key=lambda r:r['time_seconds'])
            assert len(records)==(5 if dev=='dev1' else 4)
            data=[{**r,'report_label':f"{r['frame_id']}   GT | E0 complete S1 | Wu 4DGS"} for r in records]
            picture=self.contact(dev+'_all_E_full',data)
            self.image(picture,'每行依次为真实图像 完整 S1 适配 4DGS。保留全部固定 E，完整视野没有按模型输出裁切；原高分辨率证据内嵌，可放大检查。',max_height=9.2)
        self.page('全部保留帧的固定物体裁剪')
        records=sorted([r for r in figures if r['group']=='camera1_E'],key=lambda r:(r['dev'],r['time_seconds']))
        pic=self.contact('all_E_fixed_crops',[{**r,'report_label':f"{r['dev']}  {r['frame_id']}  GT | S1 | 4DGS"} for r in records],columns=2,mode='behave_crop')
        self.image(pic,'全部九帧固定裁剪。每格依次为真实图像 完整 S1 适配 4DGS；边界由原 O 外接框增加 60 像素得到，各方法使用同框。裁剪只帮助观察，主指标仍使用整图上的原固定区域。',max_height=7.9)
        rows=[['事件帧','固定裁剪左 上 右 下']]
        for r in records:rows.append([DEV_NAMES[r['dev']]+' '+r['frame_id'],'  '.join(map(str,r['crop_bounds_xyxy']))])
        self.table(rows,[2.3,4.69],8.2)

    def fit_figures(self,figures):
        self.page('输入相机的固定规则示例')
        for group,title in [('camera0_full_training_fit','原训练帧示例'),('camera0_paired_E','同 E 名义时刻示例')]:
            selected=[]
            for dev in ['dev1','dev2']:
                rows=sorted([r for r in figures if r['dev']==dev and r['group']==group],key=lambda r:r['time_seconds'])
                assert rows
                # Selection depends on index/time only, never model scores.
                for i in sorted({0,len(rows)//2,len(rows)-1}):
                    r=rows[i];selected.append({**r,'report_label':f"{dev} {r['frame_id']}  GT | S1 | 4DGS"})
            pic=self.contact(group+'_six_examples',selected,columns=2)
            self.image(pic,title+'。每事件按已固定预览列表的首 中 末索引选三帧，共六帧，不按质量选图；列顺序始终是真实图像 完整 S1 适配 4DGS。',max_height=3.8)
        self.p('全部原训练 114 和 98 帧均已评价；这里仅压缩展示固定例子。反馈包保留固定规则预览、全部 paired E 图集与各帧指标，本地评价目录另存全部浮点前向。对 H 与 S 的欠拟合、重新显露区域和边界缺陷，应同时看完整图与物体局部，避免把背景收益直接称作交互改善。')

    def hos_results(self,H):
        self.page('移动单目 Backpack 外部校准')
        self.paras('hos_interpretation',True)
        for key,title in [('H0_test','H0 官方检查点原生条件'),('H1_test','H1 训练帧初始化的适配 4DGS'),('H1_input_fit','H1 原训练全集拟合')]:
            self.doc.add_heading(title,2)
            rows=[['区域','有效／总帧','均值 PSNR','pooled PSNR','SSIM','LPIPS']]
            for region in ['full','foreground','background']:
                d=H['results'][key][region]
                rows.append([REGION_NAMES[region] if region!='foreground' else '合并前景',f"{d['valid_frames']}/{d['frames']}",numeric(d['psnr_db']),numeric(d['pooled_psnr_db']),numeric(d['ssim'],4),numeric(d['lpips_spatial_mean'],4)])
            self.table(rows,[1.1,.85,1.25,1.35,1.2,1.24],8.4)
        self.p('三张表分开呈现，不计算 H1 减 H0 的公平增益。有效帧指该区域非空且有误差记录的帧，空区域不参与均值或 pooled 统计；NA 表示不可用，不是零误差。前景仅为发布 soft mask 以 128 为阈值的合并区域，没有独立的人 物分区。统一指标采用 AlexNet 空间 LPIPS；H0 原生程序的 VGG 标量 LPIPS 与这里不能混写。','Caption')
        if not H['lpips']['available']:
            self.p('统一 LPIPS 本次全部记 NA，原因由评价程序记录为 '+H['lpips'].get('reason','未提供具体原因')+'；这不影响已有 PSNR 与 SSIM 的有效帧。','Caption')
        native=H['native_H0_metrics']
        self.p(f"H0 原生 16 帧复核另得 PSNR {numeric(native['PSNR_frame_mean'])}、常规二维 SSIM {numeric(native['SSIM_conventional_2D_frame_mean'],4)}、VGG LPIPS {numeric(native['LPIPS_VGG_scalar_frame_mean'],4)}。它们仅标为本机原生检查点渲染指标；历史训练划分身份未完全核实，不宣称精确复现论文数值。")
        edge=H['time_boundary_diagnostic']
        self.p(f"首个测试帧 {edge['frame_id']} 位于 H1 训练帧号范围之外，归一化时间为 {edge['H1_normalized_time']:.8f}，明确属于外推并照常保留；没有静默截断到零。H0 按原生时间接口处理。移动相机测试同时涉及时间与视角变化，与 BEHAVE 同步多相机压力测试是不同问题。")

    def hos_figures(self,H):
        figures=sorted(H['figures'],key=lambda r:Path(r['path']).stem)
        assert len(figures)==16
        metrics=csv_rows(self.run/'evaluation/hos_comparison/metrics_per_frame.csv')
        lookup={(r['frame_id'],r['method'],r['region']):r for r in metrics}
        for chunk_no,start in enumerate([0,8],1):
            self.page('Backpack 全部测试图像 '+str(chunk_no))
            subset=figures[start:start+8]
            rec=[{**f,'frame_id':Path(f['path']).stem,'report_label':Path(f['path']).stem+'   GT | H0 | H1'} for f in subset]
            pic=self.contact('hos_all_test_'+str(chunk_no),rec,columns=2,mode='hos_full')
            self.image(pic,'固定测试列表按帧号排序，每格依次为真实图像 H0 官方检查点 H1 适配 4DGS。完整图像内嵌，不挑选成功帧；两个系统的信息条件不同，图中并列用于外部校准。',max_height=5.8)
            rows=[['帧 ID','H0 PSNR','H1 PSNR','H0 SSIM','H1 SSIM','H0 LPIPS','H1 LPIPS']]
            for f in rec:
                fid=f['frame_id'];a=lookup[fid,'H0','full'];b=lookup[fid,'H1','full']
                rows.append([fid,numeric(a['psnr_db']),numeric(b['psnr_db']),numeric(a['ssim'],4),numeric(b['ssim'],4),numeric(a['lpips_spatial_mean'],4),numeric(b['lpips_spatial_mean'],4)])
            self.table(rows,[.7,1.05,1.05,1.05,1.05,1.045,1.045],8.2)
            self.p('表为统一评价尺寸上的完整图逐帧指标，独立保存每个系统结果。这里没有按结果挑终态、调整门槛、换初始化或删除困难帧。','Caption')

    def costs_page(self):
        self.page('实际训练成本与完整性')
        ledger=read(self.run/'protocol/gpu_cost_ledger.json')
        rows=[['运行','名义步与更新','训练进程秒','最终与峰值点数','峰值 allocated GiB','检查点 MiB']]
        formal=[]
        for name,label in [('behave_dev1_formal','箱体 E1'),('behave_dev2_formal','木椅 E2'),('hos_backpack_formal','Backpack H1')]:
            d=read(self.run/'runs'/name/'run.json');assert d['status']=='completed'
            assert d['nominal_iterations']==17000 and d['optimizer_updates']==16999
            formal.append(d)
            rows.append([label,f"{d['nominal_iterations']}\n{d['optimizer_updates']}",numeric(d['seconds'],1),f"{d['final_points']:,}\n{d['peak_points']:,}",numeric(d['peak_allocated_bytes']/2**30,3),numeric(Path(d['checkpoint']).stat().st_size/2**20,1)])
        self.table(rows,[1.15,1.1,1.15,1.55,1.15,.89],8.2)
        temp=self.run/'protocol/temporary_steps.jsonl';temp_steps=len(temp.read_text().splitlines())
        self.p(f"三个正式运行共 {sum(r['nominal_iterations'] for r in formal):,} 个名义迭代、{sum(r['optimizer_updates'] for r in formal):,} 次优化器更新。锁定官方 coarse 3000 和 fine 14000 日程，保留官方最后 fine 步只反传而不更新的行为，因此每次 17000 对应 16999。临时检查实际 {temp_steps} 步，上限 200；临时产物不进入正式初始化。固定种子 12345，不按保留集表现补种子或选中间检查点。")
        total=sum(x['wall_seconds'] for x in ledger);peak=max((x.get('peak_process_nvidia_MiB') or 0 for x in ledger),default=0)
        self.p(f"本轮统一使用物理 GPU1 RTX 3090。运行台账的 GPU 任务进程用时合计 {total:.1f} 秒，即 {total/3600:.3f} 小时，预算为 12 GPU 小时；包含台账内训练、临时检查和前向任务的加载等工作，不是纯 CUDA 核计时。台账采样的单进程显存峰值为 {peak:,} MiB。表中 PyTorch allocated 不是整卡占用，采样峰值也可能漏掉瞬时极值。")
        if self.costs.get('preprocessing'):
            rows=[['预处理或验收','时间秒','设备','说明']]
            for p in self.costs['preprocessing']:
                rows.append([p['name'],numeric(p.get('seconds'),2),p.get('device','NA'),p.get('note','')])
            self.table(rows,[1.5,.9,1.0,3.59],8.1,left_columns=(0,3))
        for x in self.costs.get('totals',[]):self.p(f"{x['label']}：{x['value']} {x.get('unit','')}。{x.get('note','')}")
        for note in self.costs.get('notes',[]):self.p(note)
        self.paras('incidents')
        self.p('每个数据协议先冻结计划终态及哈希，再前向导出和统一评价；训练没有加载保留 RGB。旧模型、失败输出和源缓存保留。技术性失败及修复按实际日志披露；任何预算截断都不能自动视作完整方法失败。')

    def decision_and_sources(self):
        self.page('阶段决定与复算来源')
        self.paras('next_decision',True)
        self.paras('limitations',True)
        self.p('本轮不自动启动第四次正式训练，不继续注意力颜色扫参、位姿局部修补或扩大数据矩阵。后续只围绕证据支持的一个主要问题另定最小可证伪对照与预算。11 月 4 日核心结果冻结、11 月 5 日至 15 日连续 11 天集中写作窗口保持。')
        self.doc.add_heading('代码与模型身份',2)
        cfg=read(self.run/'code/experiment_config.json')
        self.p('Wu 4DGS 官方代码 commit  '+cfg['official_commit'],'Caption')
        launch=read(self.run/'protocol/launcher.json')
        self.p('本项目正式启动 commit  '+launch['project_commit'],'Caption')
        for name,label in [('behave_dev1_formal','E1'),('behave_dev2_formal','E2'),('hos_backpack_formal','H1')]:
            d=read(self.run/'runs'/name/'run.json')
            self.p(label+'  '+Path(d['checkpoint']).name+'\nSHA256  '+d['checkpoint_sha256'],'Caption')
        for dev in ['dev1','dev2']:
            s=read(self.run/'protocol'/f'behave_{dev}_native_render_check.json')['source_identity']['checkpoint']
            self.p('E0 '+DEV_NAMES[dev]+'  SHA256  '+s['sha256'],'Caption')
        h0=read(self.run/'protocol/hos_checkpoint_identity.json')
        self.p('H0 Backpack 官方检查点  SHA256  '+h0['sha256'],'Caption')
        self.p('完整输入与模型来源、初始化哈希、时间/相机清单、最终配置和成本保存在协议索引及反馈包。DOCX 已内嵌图像，复制文档后无需访问本机图片路径；代码、JSON、CSV 和 Markdown 作为可复算源保留。数据、检查点、原始图像和第三方源码不进入公开代码同步。','Caption')
        self.doc.add_heading('主要资料与复算文件',2)
        for src in self.content.get('sources',[]):
            self.p(src['label']+'  '+src.get('url','')+'\n'+src.get('detail',''),'Caption')
        self.p('本轮数值源  existing_error_summary.json  evaluation/comparison/summary.json  evaluation/hos_comparison/summary.json  对应 metrics_per_frame.csv 和 input_fit.csv  protocol/gpu_cost_ledger.json  costs.json。正式方法适配入口为 code/adapter_4dgs.py 与 code/train_official.py，完整旧 S1 恢复入口为 code/render_full_s1.py。','Caption')


def main(args):
    run=args.run.resolve();content=read(args.content or run/'report_content.json');costpath=args.costs or run/'costs.json';costs=read(costpath)
    for key in ['summary_paragraphs','behave_interpretation','hos_interpretation','next_decision','limitations']:
        assert content.get(key) and isinstance(content[key],list),f'Missing root-authored {key}'
    B=read(run/'evaluation/comparison/summary.json');H=read(run/'evaluation/hos_comparison/summary.json');A=read(run/'existing_error_summary.json')
    fmanifest=read(run/'evaluation/comparison/figure_manifest.json');figures=fmanifest['figures']
    coverage=validate_coverage(run,B,H,fmanifest)
    r=Report(run,content,costs)
    r.summary(B,A);r.protocol();r.old_errors(A);r.main_regions(B);r.fit_and_pairs(B);r.behave_figures(figures);r.fit_figures(figures);r.hos_results(H);r.hos_figures(H);r.costs_page();r.decision_and_sources()
    expected_e={(dev,fid) for dev,ids in coverage['behave_E_ids'].items() for fid in ids}
    for mode in ['behave_full','behave_crop']:
        included=[(x['dev'],x['frame_id']) for x in r.image_sources if x['mode']==mode and x['group']=='camera1_E']
        assert len(included)==9 and set(included)==expected_e,'DOCX must embed all nine E full images and fixed crops'
    included_hos=[x['frame_id'] for x in r.image_sources if x['mode']=='hos_full']
    assert len(included_hos)==16 and set(included_hos)==set(coverage['hos_test_ids']),'DOCX must embed all sixteen HOS test images'
    output=args.output or run/'output/V3_baseline_calibration.docx';output.parent.mkdir(parents=True,exist_ok=True);r.doc.save(output)
    with zipfile.ZipFile(output) as z:
        media=[name for name in z.namelist() if name.startswith('word/media/')]
        rels=z.read('word/_rels/document.xml.rels').decode()
        assert 'TargetMode="External"' not in rels,'Report images must be embedded, no external image dependency'
    audit=dict(status='authored_awaiting_render_and_visual_QA',created_utc=datetime.now(timezone.utc).isoformat(),docx=ident(output),
        planned_pages=14,sections=r.section_pages,embedded_media_count=len(media),image_sources=r.image_sources,figure_coverage=coverage,
        sources=[ident(p) for p in [args.content or run/'report_content.json',costpath,run/'existing_error_summary.json',
            run/'evaluation/comparison/summary.json',run/'evaluation/comparison/figure_manifest.json',run/'evaluation/hos_comparison/summary.json',
            run/'evaluation/comparison/metrics_per_frame.csv',run/'evaluation/comparison/input_fit.csv',
            run/'evaluation/hos_comparison/metrics_per_frame.csv',run/'evaluation/hos_comparison/input_fit.csv',Path(__file__)]],
        behavior='Saved metrics and precomputed comparison figures only; no model runs, metric recomputation, new GT decoding or optimization',
        formal_nominal_steps=51000,formal_optimizer_updates=50997,
        skill_mark_executed_by='caller exactly once before first build',visual_QA_executed=False)
    (run/'output/report_build_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in audit.items() if k not in ['image_sources','sources','sections']},ensure_ascii=False,indent=2))


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__,epilog='After DOCX render and visual QA, use code/package_results.py for the feedback archive.');p.add_argument('--run',type=Path,default=RUN);p.add_argument('--content',type=Path);p.add_argument('--costs',type=Path);p.add_argument('--output',type=Path)
    main(p.parse_args())
