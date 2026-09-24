"""Portable concise report with embedded figures; run with bundled document Python."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
E=Path(__file__).resolve().parents[1]
syn=json.loads((E/'synthetic_contract/results.json').read_text());real=json.loads((E/'mask_ledger/results.json').read_text());audit=json.loads((E/'protocol/audit_result.json').read_text())
D=Document();s=D.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=D.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');st.paragraph_format.space_after=Pt(6)
    grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(grid)
D.styles['Normal'].font.size=Pt(10.5);D.styles['Normal'].paragraph_format.line_spacing=1.13
D.styles['Title'].font.size=Pt(20);D.styles['Heading 1'].font.size=Pt(15);D.styles['Heading 2'].font.size=Pt(12);D.styles['Caption'].font.size=Pt(9)
D.styles['Caption'].font.bold=False;D.styles['Caption'].font.italic=False;D.styles['Subtitle'].font.italic=False
for st in D.styles:
    for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
def p(t,style=None):return D.add_paragraph(t,style)
def page(t):
    q=D.add_heading(t,level=1);q.paragraph_format.page_break_before=True
def table(rows,widths):
    t=D.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,(cell,txt,w) in enumerate(zip(cells,row,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for tag,value in [('top',75),('bottom',75),('left',85),('right',85)]:
                el=OxmlElement('w:'+tag);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');m.append(el)
            pr.append(m);b=OxmlElement('w:tcBorders')
            for tag in ['top','left','bottom','right']:
                el=OxmlElement('w:'+tag);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            pr.append(b)
            if i==0:
                el=OxmlElement('w:shd');el.set(qn('w:fill'),'E8EEF4');pr.append(el)
            q=cell.paragraphs[0];q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=1.05;q.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            r=q.add_run(str(txt));r.font.size=Pt(9.5);r.bold=i==0
        if i==0:
            el=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(el)
        el=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(el)
    p('')
def image(path,width):
    q=p('');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(3)
    shape=q.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set('descr',path.stem.replace('_',' '))
D.core_properties.title='独立位姿轮廓监督核对与研究收敛';D.core_properties.author='HOI research project'
p('独立位姿轮廓监督核对与研究收敛','Title');p('2026年9月24日  更新附件第14节执行结果','Subtitle')
p('结论：轮廓语义错误假设未成立。原求解器已区分可见物体、可见人体和背景；本轮完成源码核对、一个人工遮挡例与六个输入选定真实帧检查。按附件停止条件，新增位姿求解0次、高斯训练0次、独立事件0条，保留旧S1及全部历史失败。')
p('要检查的问题是：单独投影的完整物体可能藏在人体后面，若把人体截断后的可见物体边缘当作完整外边界，就可能错误收缩或移动物体。实际代码的背景距离场只在输入背景B内为正，在可见物体O、人体H内为0；两缓存没有显式不确定区U，本轮不人为新增。')
p('箱体正向项要求64个O像素落在完整模板投影凸包内，内部距离为0；木椅使用O像素到完整投影表面采样点的最近距离。后者存在采样近似，但没有要求隐藏轮廓贴合人体切口。原采样、16 px尺度、完整帧数分母和固定可见帧均保留。')
image(E/'synthetic_contract/contract.png',7.05)
p('同一个人工场景。黑线为预测完整轮廓，青色为可见物体O，粉色为人体H，紫色为人工U，灰色为背景B。三列依次为完整预测、删除可见前景、侵入背景。该例仅检验代码，不是重建性能结果。','Caption')
table([['人工条件','箱体正向','箱体背景','木椅正向','木椅背景'],*[ [name,*[f'{syn["cases"][key][dev][part]:.3f}' for dev in ['dev1','dev2'] for part in ['silhouette_forward','silhouette_background']]] for key,name in [('complete','完整投影'),('delete_visible','删除可见前景'),('invade_background','侵入可信背景')] ]],[1.85,1.3,1.3,1.3,1.3])
p('H/U内部探针的背景残差与梯度均为0；隐藏边在H内移动，正向和背景值均不变。删除前景的梯度指向恢复覆盖；侵入背景受罚。木椅完整预测的0.234来自2 px网格采样近似。旧CPU与现行Torch分项差为0，原求解器未修改。','Caption')

for dev,name in [('dev1','箱体'),('dev2','木椅')]:
    page(name+'真实输入中的监督作用')
    p('每序列按时间分三段，每段选择O与H相邻占比最高、且同时有O与B边界的固定可见帧；规则在查看残差前冻结，不按参考或模型误差挑选。黄色为历史F01完整模板投影，彩色查询显示原始像素距离。')
    image(E/f'mask_ledger/{dev}_frozen_frames.png',6.4)
    data=[['帧','O像素','H内部采样','H残差最大','H梯度最大','B内部采样']]
    for row in real[dev]['frames']:
        reg=row['regions'];data.append([row['frame'],row['pixel_counts']['O'],reg['H']['count'],reg['H']['max_raw_px'],reg['H']['max_raw_gradient_norm'],reg['B']['count']])
    table(data,[.6,1.05,1.45,1.3,1.3,1.35])
    if dev=='dev1':p('H内部指双线性插值的四个像素角都属于人体。箱体三帧共95个这样的实际查询，背景距离与梯度均为0；原O全部保留。按最近像素分类，另有1个H边界查询产生0.334 px小残差，来源是插值单元包含背景像素。','Caption')
    else:p('木椅三帧共3078个H内部查询，背景距离与梯度均为0。另有42个最近像素为H的边界查询非零，最大0.813 px，均可由含B的亚像素插值解释。该边界带不构成把整个人体当背景的证据；未扩张忽略区。','Caption')

page('当前研究范围收敛决定')
p('本轮于9月24日按附件第8.5与14节同日收口。人工例及真实输入没有建立具体可修正的遮挡语义冲突，因此不生成C1，不重放C0，不继续局部位姿修补。F01保持失败对照身份；没有新的配对三维、高斯、camera1或接触结论。')
D.add_heading('已有证据支持停止的动作',level=2)
p('单源固定对应退化；P1及S1*只在木椅绝对位置上部分有效，跨事件运动没有统一改善；关闭短轨迹没有纠正箱体漂移；官方同池重排没有统一收益。F01仅相对退化F00有有限运动作用，可调q及联合F11均没有升级依据。本轮轮廓语义错误假设也未成立。')
p('保留旧S1和失败记录，不重复八臂、扩大5 cm邻域、换特征或扫权重与mask宽度；不自动追加关键帧硬锚、全局朝向搜索、注意力、接触或新网络。')
D.add_heading('仍缺证据与主张边界',level=2)
p('在标定单目RGB、RGB预测人体姿态及已知米制物体形状的协议下，尚未建立在两dev统一优于“不更新旧初始化”的可靠物体运动纠正机制。二维损失下降不保证深度与运动正确；标签、弱深度、模板投影近似及局部优化仍可能影响结果，本轮没有锁定唯一原因。')
p('原代码可运行、工程可复算，与论文核心机制已有证据应明确区分。可信材料对应、可用粗运动后实际局部高斯支持丢失、固定规则确实误配等后续启用证据仍不足。当前不能主张稳定运动恢复、精细接触改善或已有充分CVPR贡献，也不能将有限失败写成数学上不可恢复。')
D.add_heading('下一阶段边界与复算',level=2)
p('不消耗八条尚未调参事件寻找容易成功例。最迟10月7日收敛一个有证据的核心问题；若更换任务、输入或机制，先明确协议、独立增量、最小支持与否定实验及停止条件。11月4日冻结核心结果，11月5—15日连续11个完整写作日不变。')
p(f'本轮数值核对与原始出图为CPU执行，实测{audit["wall_seconds"]:.2f}秒，不含阅读、报告编写及版式检查；未占用GPU，未创建后台训练。原代码、标签、历史产物哈希另行核验。','Caption')
p('实验根：experiments/visibility_objective_check_20260924/run01。TARGET_SEMANTICS.md包含函数行号与限制；protocol固定输入及选帧；synthetic_contract和mask_ledger保存分项、梯度与数组；REPRODUCE.md列出命令；RESEARCH_SCOPE_DECISION.md保存本页决定。DOCX的三张图全部内嵌。','Caption')
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
out=E/'output/Visibility_objective_check.docx';D.save(out);print(out)
