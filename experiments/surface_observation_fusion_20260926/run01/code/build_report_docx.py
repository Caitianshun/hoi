"""V2 report. Uses saved results only; bundled DOCX runtime."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
E=Path(__file__).resolve().parents[1];R=E.parents[2];B=R/'experiments/fixed_motion_reconstruction_20260926/run01'
def read(p):return json.loads(p.read_text())
S=read(E/'evaluation/summary.json');rows=read(E/'evaluation/per_frame.json');I=read(E/'protocol/final_integrity.json');F=read(E/'protocol/fusion_frozen.json');REP=read(E/'evaluation/representation.json')
D=Document();sec=D.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.7)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=D.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');st.paragraph_format.space_after=Pt(6)
D.styles['Normal'].font.size=Pt(10.5);D.styles['Normal'].paragraph_format.line_spacing=1.1;D.styles['Title'].font.size=Pt(19);D.styles['Heading 1'].font.size=Pt(15);D.styles['Heading 2'].font.size=Pt(12);D.styles['Caption'].font.size=Pt(9);D.styles['Caption'].font.italic=False
for st in D.styles:
 for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
def p(t,style=None):return D.add_paragraph(t,style)
def page(t):q=D.add_heading(t,1);q.paragraph_format.page_break_before=True
def table(data,widths,size=9):
 t=D.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,row in enumerate(data):
  cells=t.add_row().cells
  for c,text,w in zip(cells,row,widths):
   c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for tag,value in [('top',65),('bottom',65),('left',70),('right',70)]:el=OxmlElement('w:'+tag);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar);b=OxmlElement('w:tcBorders')
   for tag in ['top','left','bottom','right']:el=OxmlElement('w:'+tag);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
   pr.append(b)
   if i==0:el=OxmlElement('w:shd');el.set(qn('w:fill'),'E8EEF4');pr.append(el)
   q=c.paragraphs[0];q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=1.0;q.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=q.add_run(str(text));r.font.size=Pt(size);r.bold=i==0
  el=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(el)
  if i==0:el=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(el)
 p('').paragraph_format.space_after=Pt(0)
def image(path,width=7):
 q=p('');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(3);q.add_run().add_picture(str(path),width=Inches(width))
def fmt(x,digits=3):return 'NA' if x is None else f'{x:.{digits}f}'
p('固定运动监督与表面观测融合验证','Title');p('2026年9月26日  V2 执行总结  两条既有开发序列','Subtitle')
p('六次正式优化与统一评价已完成。B1 独立可见监督没有达到两事件各 +0.5 dB 的工程门槛；F2 注意力相对同信息 MLP 的增量也远小于 +0.2 dB。当前证据支持保留缓存与有限局部改善记录，不支持升级 B1 或宣称注意力必要。本轮收口，不在两事件继续扫参或追加训练。')
p('这轮区分两个问题：物体单独渲染的可见区域监督是否有帮助；在固定几何上选择多时刻真实 RGB，是否优于固定加权以及普通打分器。B0 是旧 AUX 的 RR 终态；B1 从同一个未训练物体初值优化。F0/F1/F2 全部挂在冻结 B0 上，不能把它们与 B1 的差异解释为单一算子因果贡献。')
t=[['事件','版本','O PSNR dB','O SSIM','O LPIPS']]
for d,n in [('dev1','箱体'),('dev2','木椅')]:
 for v in ['B0','B1','F0','F1','F2']:
  m=S[d]['variants'][v]['object'];t.append([n,v,fmt(m['psnr_db']['mean']),fmt(m['ssim']['mean'],4),fmt(m['lpips_spatial_mean']['mean'],4)])
table(t,[1,1,1.6,1.5,1.9])
p('表为每帧固定物体 O 区域等权均值。PSNR 衡量像素误差，SSIM 衡量局部结构，LPIPS 是感知差异；前两者越大越好，LPIPS 越小越好。完整区域、不确定区域与箱体 t26 的 200 像素均保留。SSIM 窗口及 LPIPS 感受野跨边界，不是纯表面误差。','Caption')
p('B1 子协议为 AUX_FIXED_MOTION_RECON_V1；F 系列为 AUX_SURFACE_OBSERVATION_FUSION_V1。只使用原箱体 12/5 帧、木椅 9/4 帧训练/评价集合。Ref 物体拟合运动是额外输入，并非完美真值或自动 RGB 恢复；人体、背景与人体运动取旧 S1 并冻结。camera1、发布人体、深度和纹理不进入训练。')
page('实现机制与预声明判断')
p('B1 每步同时渲染完整场景和黑背景物体。独立颜色项只在原 O 内部要求解释 RGB，在背景内部抑制溢出；独立 alpha 项相应要求 O 有支撑、背景无物体。人体及边界不取独立损失。两个项权重为 0.5 和 0.1；原合成损失与静态正则保持。两路物体屏幕梯度向量先相加，再做原增密统计，seen_count 每步只计一次并集。故 B1 同时改变属性更新和普通增长过程。')
p('融合先用物体三角网格首表面把 camera0 像素映射回规范坐标，保留 5×5 腐蚀后物体内部的真实观测。每个实际高斯中心只找继承面上的 5 mm 局部，每个源时刻一条，最多八条。面关联不唯一或缺少足够来源时回退 B0；子代按自己的实际位置查询，不复制锚点平均色。')
p('固定可靠性是内部距离代理乘 5 mm 尺度的高斯距离衰减。F0 用其归一化权重；F1 用同时读取查询与候选描述的两层 MLP 修正分数；F2 用单头 32 维 QK 点积修正。查询有位置、法向、B0 颜色与候选中位描述，候选有真实 RGB、有效 patch、梯度、方向等；没有可学习锚点 ID、自由颜色残差或新网络编码器。')
p('监督时先排除当前来源，再选候选并重算所有聚合；不足两个来源共同回退。B0 本身见过完整 camera0，因此这不是严格未见帧测试。F1/F2 仅训练共享打分参数，均由 F0 输出起步；F1 为 4096 参数，F2 为 4064 参数，未设置 softmax 会抵消的输出或 key 偏置。各自 2000 步 Adam，lr=0.001，末态固定；全缓存烘焙一次静态颜色后才评价。')
t=[['配对','箱体均值 / 中位数','木椅均值 / 中位数']]
for pair in ['B1-B0','F0-B0','F1-F0','F2-F0','F2-F1','F2-B0','F2-B1']:
 vals=[S[d]['differences'][pair]['object']['psnr_db'] for d in ['dev1','dev2']];t.append([pair]+[f"{v['mean']:+.3f} / {v['median']:+.3f}" for v in vals])
table(t,[1.2,2.9,2.9])
p('B1 的均值增益未达 +0.5 dB；其余 SSIM、完整图、贡献比例、支持区域及边界护栏通过。F2−F1 未达 +0.2 dB；木椅逐帧差中位数还为负。F2 相对 F0/B0 均值未下降，其他数值护栏通过，但不足以保留本次注意力作为有效增量。这些是资源分配工程标准，不是统计显著性或论文创新标准。')
for d,n in [('dev1','箱体'),('dev2','木椅')]:
 page(n+'全部评价帧')
 image(E/f'output/figures/{d}_crop.png',6.8)
 p('每行依次为 GT、B0、B1、F0、F1、F2、仅 H/S。青线为同一固定 O；裁剪由参考 O 框加 60 像素决定。全部帧都展示，不按收益挑选。','Caption')
 rr=sorted([r for r in rows if r['dev']==d and r['variant']=='B0'],key=lambda r:r['time']);lookup={(r['time'],r['variant']):r for r in rows if r['dev']==d};t=[['时刻','B0','B1','F0','F1','F2']]
 for r in rr:
  cells=[f"t{r['time']:g}"]
  for v in ['B0','B1','F0','F1','F2']:
   m=lookup[r['time'],v]['metrics']['object'];cells.append(f"{m['psnr_db']:.3f}\n{m['ssim']:.4f}\n{m['lpips_spatial_mean']:.4f}")
  t.append(cells)
 table(t,[.65,1.27,1.27,1.27,1.27,1.27],8)
 p('每单元三行依次为 PSNR dB、SSIM、空间 LPIPS。模型输出保存为浮点数组，表格不是从展示用 PNG 反算。','Caption')
page('输入支持覆盖与表示限制')
p('支持标签描述的是给定 Ref 几何和输入 mask 下的保守代理：至少两个原生时刻为 supported，一次为 single_view；有局部几何映射但没有 O 正证据为 no_positive_evidence；没有可靠面关联或 5 mm 局部采样为 unmapped。无证据可能是分割或采样限制，不能直接称真实表面不可见。')
t=[['事件时刻','O 像素','至少两源','一次','无正证据','未映射','至少三源']]
for d,n in [('dev1','箱'),('dev2','椅')]:
 for x in sorted((E/'evaluation'/d).glob('*/support.json')):
  j=read(x);c=j['counts'];den=j['O_pixels'];t.append([n+f" t{j['time']:g}",den]+[f"{c[k]}\n{100*c[k]/den:.1f}%" for k in ['supported','single_view','no_positive_evidence','unmapped','source_ge3']])
table(t,[1,.65,1.05,1.05,1.05,1.05,1.15],8.5)
p('全部 9 帧固定 O 均有模板深度；这不等于模板或拟合正确。未映射比例单列，不排除出主指标。各帧 supported 均不少于 100 像素，因此 B1 支持区域护栏在两事件都启用。箱体/木椅 supported PSNR 配对均值为 +0.435/−0.005 dB，均过 −0.2 dB 护栏。','Caption')
p('B0 的实际高斯中，箱体 2782/6000、木椅 1396/6000 个具有至少两源；至少三源为 1564/6000、470/6000。更严格的继承面唯一关联使 1011/1628 个高斯不能可靠映射。表面有正观测却在同面 5 mm 内没有初始 anchor 的评价像素也已保存，不能归为未观测。这是容量或采样的后续线索，不是已证实的主要原因；当前高斯核仍可能覆盖邻近像素。')
p('箱体 F2 相对 F0 的 supported PSNR 提高约 0.208 dB，而木椅约 0.024 dB。收益并未统一转化为完整 O 的明显改进。固定几何、opacity 与 H/S 的误差不会由颜色选择修复；只靠本轮结果不能区分对应误差、采样限制和表示上限的相对贡献。')
page('来源审计与回退记录')
p('每事件按实际高斯局部 ID 等间距取八个样本，含无来源和单来源回退。图中为源像素周围上下文，青色十字是实际读色中心；上下文图片可能出现其他实体，但输入颜色和描述有效像素均按原规则过滤。t 是原生时刻，w 是 F2 最终权重。')
image(E/'output/figures/all_sources.png',6.5)
p('箱体 ID2571 在 t23/27/29/30 权重约 0.18/0.66/0.10/0.06；其邻域有明显亮度和边界变化。木椅 ID0 在 t4/5/6 为 0.56/0.28/0.16。规则确保同面、局部距离和源标签符合输入条件，但这些并非材料对应真值，不能从权重图断言选对纹理。其余样本及可靠性、源 RGB、坐标、回退原因均保存在 source_audit.json。','Caption')
p('视觉核对未见代码跨实体取中心像素的违规；几何偏差仍可能把同一面上不同纹理混在 5 mm 邻域内，尤其是印字、椅背边缘。此为可见风险而非新增真值误配率。单凭 mask 消失不能确认真实遮挡，因此可见—遮挡—再显露专项测试记为 NA，不主张长时间遮挡恢复。','Caption')
page('全部评价帧的完整组合图')
image(E/'output/figures/dev1_full.png',7)
image(E/'output/figures/dev2_full.png',7)
p('上为箱体全部五帧，下为木椅全部四帧；列顺序 GT、B0、B1、F0、F1、F2、H/S-only。完整图内嵌高分辨率源，可放大查看。冻结背景和人体在新视角仍有明显缺失与模糊；各版本都共享这些问题。完整图 PSNR 约 7.3–7.4 dB，不能将小范围物体颜色变化写成整体三维恢复。','Caption')
page('表示完整性和实际成本')
t=[['事件版本','点数','最大尺度 mm','最大偏移 mm','opacity 中位数']]
for d,n in [('dev1','箱体'),('dev2','木椅')]:
 for v in ['B0','B1']:
  a=REP[d+'_'+v];t.append([n+' '+v,a['gaussian_count'],fmt(a['scale_each_axis_m']['max']*1000,2),fmt(a['canonical_offset_norm_m']['max']*1000),fmt(a['opacity']['median'])])
table(t,[1.4,.7,1.7,1.6,1.6])
p('F0/F1/F2 除烘焙颜色外，B0 的中心、尺度、方向、opacity、拓扑与缓冲逐张量一致；运动和 H/S 始终冻结。B1 保持同一 6000 点上限及 5 mm 偏移边界。只固定中心不等于约束高斯核厚度，35 mm 仍是旧尺度软正则起点。')
t=[['事件分支','名义 / 有效更新','进程秒','峰值 allocated GiB']]
for d,n in [('dev1','箱体'),('dev2','木椅')]:
 for v in ['B1','F1','F2']:
  j=read((B/'runs'/d if v=='B1' else E/'runs'/f'{d}_{v}')/'run.json');t.append([n+' '+v,f"{j['steps']} / {j.get('updates',j['steps'])}",fmt(j['seconds'],2),fmt(j['peak_allocated_bytes']/2**30)])
table(t,[1.4,2.1,1.5,2.0])
p(f"本机物理 GPU1 RTX3090，六次优化合计 24000 步，六次临时检查合计 48 步。GPU 使用进程用时总和 {I['GPU_process_seconds_sum']:.2f} 秒，含加载等 CPU 工作，不是 CUDA 核计时；两条流水线及首次检查合计 {I['pipeline_wall_upper_bound_seconds']:.2f} 秒，远低于 4 GPU 小时预算。峰值 allocated {I['max_peak_allocated_bytes']/2**30:.3f} GiB，不是整卡显存。缓存构建约 9.29 秒，CPU 特征与输入核对约 4.16 秒；缓存压缩资产 {I['cache_bytes']/2**20:.2f} MiB。历史人体估计、S1与 B0成本属于复用成本，不记作零。")
p('F0 零优化。所有 F 模型完成全缓存静态颜色烘焙；每个输入相机帧的 joint 与 object-only 导出一致性最大误差均为 1.19×10⁻⁷，低于 1×10⁻⁵。导出颜色 logit 使用统一 epsilon=10⁻⁶，本轮未触发颜色裁切。所有正式步骤都有有限非零梯度，no-op 为 0；空 O 帧仍按原日程计算完整场景及背景项。')
p('一次持久启动器错误解析 Python 符号链接，误入缺少 smplx 的环境，发生在导入阶段，未产生预检查或正式步。修正执行路径后原计划全部完成；错误日志保留，无换种子或重训。临时参数不进入正式训练。')
page('阶段决定与复算交付')
p('H-SUP：独立物体监督有小幅外观变化，但两事件均未过 +0.5 dB，不统一升级 B1。H-FUSE：确有真实多时刻输入、真实渲染梯度和有限外观改善，属于有效可测对照；F2 与同信息 F1 的差距很小，且没有统一方向，不保留本次 attention 参数化作为必要增量。F0 及学习版的增益均弱，不能据此宣布当前融合已成为有效系统改进。')
p('本轮到此结束，保留全部六次终态、F0、新缓存和失败证据。P0 与旧 S1、旧 AUX 停止结论不变；八条未调参事件未使用。不追加种子、权重、时间窗或新模块。下一阶段若继续，先围绕“有输入证据但缺少实际表面支撑”或“固定几何限制”制定单一可证伪实验和预算；本报告不构成自动执行该后续的授权。10月7日前收敛问题，11月4日核心冻结，11月5—15日完整写作窗口保持。')
p('67 个训练前/评价前冻结资产 hash 一致；18 个原评价 RGB 与区域文件身份核对通过。45 个保存浮点渲染的完整 O PSNR/SSIM独立复算，最大差 3.55×10⁻¹⁵ / 0；所有逐帧配对差重算一致，B0 重渲染相对历史逐像素最大差为 0。LPIPS复用已安装实现及本地权重，未下载新网络，未另行重复其全部分数。')
p('代码与源配置入口：fixed_motion_reconstruction_20260926/run01/code 和 surface_observation_fusion_20260926/run01/code。输入协议、源附件、冻结索引、检查、全部逐帧 JSON/CSV、源审计、错误和实际命令保存在两实验目录的 PROTOCOL、protocol、support、runs、evaluation、REPRODUCE 与 NEXT_DECISION。完整本机索引记录模型路径及 SHA256，模型和数据不上传。')
p('GitHub 仅按项目白名单同步自产代码、源配置、README/AGENTS 和必要补丁。附件、研究日志、文档成品、图像、缓存、数据集、模型、评价输出及凭据排除；本轮未使用远端 GPU，也未向 cts 部署 HOI 数据。')
p('本报告为可移植 DOCX，图片已嵌入；另保留可复算源。仅进行本机渲染和单独文档复制验收，不冒称已在另一主机的 Word 应用逐项验收。两个既有开发事件与单种子不能支持跨场景显著性、几何恢复、接触正确或长遮挡恢复主张。','Caption')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
D.core_properties.title='固定运动监督与表面观测融合验证';D.core_properties.author='HOI research project';out=E/'output/V2_supervision_surface_fusion.docx';D.save(out);print(out)
