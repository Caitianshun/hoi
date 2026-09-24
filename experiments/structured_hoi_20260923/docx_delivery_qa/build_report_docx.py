"""Faithful Word conversion with embedded images/fonts and source coverage QA."""
from pathlib import Path
import re,json,hashlib,struct,uuid,zipfile,io
from lxml import etree
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

E=Path('/home/cai_tianshun/Project/HOI/experiments/structured_hoi_20260923')
QA=E/'docx_delivery_qa';PACKAGE=E/'portable_report';SOURCE=E/'REPORT.md'
LATIN='DejaVu Sans';CJK='Droid Sans Fallback';WIDTH=7.0
doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.65);section.bottom_margin=Inches(.65)
section.left_margin=Inches(.75);section.right_margin=Inches(.75)
section.footer_distance=Inches(.27)

def fonts(style,size,bold=False):
    style.font.name=LATIN;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for name in ['ascii','hAnsi','cs']:rf.set(qn('w:'+name),LATIN)
    rf.set(qn('w:eastAsia'),CJK)
    for a in list(rf.attrib):
        if a.endswith('Theme'):del rf.attrib[a]

for name,size,bold in [('Normal',11,False),('Title',22,True),('Heading 1',15,True),('Heading 2',12,True),('Caption',9.5,False),('List Bullet',11,False),('List Number',11,False)]:
    fonts(doc.styles[name],size,bold)
    pf=doc.styles[name].paragraph_format;pf.line_spacing=1.16;pf.space_after=Pt(7);pf.widow_control=True
for name in ['Heading 1','Heading 2']:
    pf=doc.styles[name].paragraph_format;pf.keep_with_next=True;pf.space_before=Pt(14);pf.space_after=Pt(7)
doc.styles['Title'].paragraph_format.space_after=Pt(12)
for border in list(doc.styles['Title'].element.xpath('./w:pPr/w:pBdr')):
    border.getparent().remove(border)
for name in ['Source Path','Reference Caption']:
    if name not in doc.styles:doc.styles.add_style(name,1)
fonts(doc.styles['Source Path'],8.5)
doc.styles['Source Path'].paragraph_format.space_after=Pt(4)
doc.styles['Source Path'].paragraph_format.line_spacing=1.06
fonts(doc.styles['Reference Caption'],10,True)
doc.styles['Reference Caption'].paragraph_format.keep_with_next=True
doc.styles['Reference Caption'].paragraph_format.space_after=Pt(4)

def xmlfont(r,bold=False,code=False):
    rp=r.get_or_add_rPr();rf=OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'),LATIN);rf.set(qn('w:hAnsi'),LATIN);rf.set(qn('w:eastAsia'),CJK);rp.append(rf)
    if bold:rp.append(OxmlElement('w:b'))
    if code:
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');rp.append(sz)

def hyperlink(p,label,target,bold=False):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');xmlfont(r,bold);color=OxmlElement('w:color');color.set(qn('w:val'),'1F4E79');r.get_or_add_rPr().append(color)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

link_re=re.compile(r'(?<!!)\[([^\]]+)\]\(([^)]+)\)')
text=SOURCE.read_text();references=[]
for m in link_re.finditer(text):
    if m.group(2).startswith('/') and not any(x['source']==m.group(2) for x in references):
        references.append({'id':f'R{len(references)+1:02d}','label':m.group(1),'source':m.group(2)})
refmap={x['source']:x for x in references}
mapping=json.loads((PACKAGE/'path_map.json').read_text())
if 'source_to_relative' in mapping:mapping=mapping['source_to_relative']
if 'files' in mapping and isinstance(mapping['files'],dict):mapping=mapping['files']

def relative(path):
    v=mapping[str(path)]
    return v if isinstance(v,str) else v.get('relative_path',v.get('target'))

token=re.compile(r'\*\*(.+?)\*\*|`([^`]+)`|\[([^\]]+)\]\(([^)]+)\)')
def inline(p,s,bold=False):
    pos=0
    for m in token.finditer(s):
        if m.start()>pos:p.add_run(s[pos:m.start()]).bold=bold
        if m.group(1) is not None:inline(p,m.group(1),True)
        elif m.group(2) is not None:
            r=p.add_run(m.group(2));r.font.size=Pt(9.5);r.bold=bold
        else:
            label,url=m.group(3),m.group(4)
            if url in refmap:hyperlink(p,label,relative(url),bold);p.add_run(' ['+refmap[url]['id']+']')
            else:hyperlink(p,label,url,bold)
        pos=m.end()
    if pos<len(s):p.add_run(s[pos:]).bold=bold

image_log=[];block_log=[];table_count=0
def picture(path,caption,width=WIDTH,crop=None):
    im=Image.open(path);w,h=im.size
    displayh=h if crop is None else h*(crop[1]-crop[0])
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
    shape=p.add_run().add_picture(str(path),width=Inches(width),height=Inches(width*displayh/w))
    shape._inline.docPr.set('descr',caption)
    if crop:
        fill=shape._inline.graphic.graphicData.pic.blipFill
        rect=OxmlElement('a:srcRect');rect.set('t',str(round(crop[0]*100000)));rect.set('b',str(round((1-crop[1])*100000)));rect.set('l','0');rect.set('r','0');fill.insert(1,rect)
    p=doc.add_paragraph(caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(9)
    image_log.append({'source':str(path),'sha256':hashlib.sha256(Path(path).read_bytes()).hexdigest(),'caption':caption,'crop':crop})

def set_cell_border(cell):
    tcpr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
    tcpr.append(b);mar=OxmlElement('w:tcMar')
    for edge,value in [('top','75'),('bottom','75'),('left','90'),('right','90')]:
        x=OxmlElement('w:'+edge);x.set(qn('w:w'),value);x.set(qn('w:type'),'dxa');mar.append(x)
    tcpr.append(mar)

widths=[[.7,3.1,3.2],[1.1,2.95,2.95],[.65,.6,1.35,1.35,1.525,1.525],[.65,1.7,1.55,1.9,1.2],[.65,.6,1.55,1.55,1.325,1.325],[.65,1.45,1.45,1.65,1.8],[1.65,5.35]]
def table(rows):
    global table_count
    ws=widths[table_count];table_count+=1
    t=doc.add_table(rows=len(rows),cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,ws):c.width=Inches(w)
    for i,(row,values) in enumerate(zip(t.rows,rows)):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for j,(cell,v) in enumerate(zip(row.cells,values)):
            cell.width=Inches(ws[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;set_cell_border(cell)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.10
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or (table_count in [3,4,5,6]) else WD_ALIGN_PARAGRAPH.LEFT
            inline(p,v,i==0)
            for run in p.runs:run.font.size=Pt(9.5)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DBE5F1');cell._tc.get_or_add_tcPr().append(sh)
    # Repeated header, flexible row heights, no arbitrary keep-whole-table rule.
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=text.splitlines();i=0
while i<len(lines):
    l=lines[i].strip()
    if not l:i+=1;continue
    if l.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in cells):group.append(cells)
            i+=1
        table(group);block_log.append({'kind':'table','cells':group});continue
    m=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',l)
    if m:picture(Path(m.group(2)),m.group(1));block_log.append({'kind':'image','alt':m.group(1),'path':m.group(2)});i+=1;continue
    if l.startswith('# '):p=doc.add_paragraph(style='Title');inline(p,l[2:],True);block_log.append({'kind':'heading','text':l[2:]})
    elif l.startswith('## '):p=doc.add_paragraph(style='Heading 1');inline(p,l[3:],True);block_log.append({'kind':'heading','text':l[3:]})
    else:
        body=l
        if l.startswith('- '):p=doc.add_paragraph(style='List Bullet');body=l[2:]
        else:p=doc.add_paragraph()
        inline(p,body);block_log.append({'kind':'paragraph','text':body})
    i+=1

doc.add_page_break();doc.add_heading('附录 A 固定五帧完整对照',level=1)
doc.add_paragraph('左列为输入，中列为 S0，右列为 S1。下列五行完整保留原报告所链接对照图的全部内容，按原顺序分行展示。')
jpg=E/'summary/dev2_input_S0_S1.jpg'
for i,frame in enumerate([0,24,48,72,97]):
    if i==3:doc.add_page_break();doc.add_heading('附录 A 固定五帧完整对照续',level=1)
    picture(jpg,f'对照图第 {i+1} 行  输入帧 {frame}',crop=(i/5,(i+1)/5))

doc.add_page_break();doc.add_heading('附录 B 引用资料及来源',level=1)
doc.add_paragraph('正文图片与附录 A 已嵌入本 Word 文件。视频、详细报告和机器可读记录保存在随附阅读包的 report_materials 目录；打开这些引用资料时，应一并转移阅读包。下列原路径仅作来源定位，正文阅读不依赖原主机。训练数据、模型权重及完整检查点不包含在阅读包内。')
for ref in references:
    p=doc.add_paragraph(ref['id']+'  '+ref['label'],'Reference Caption')
    p.add_run('    ')
    hyperlink(p,'打开随包资料',relative(ref['source']));p.add_run('  '+Path(ref['source']).suffix.lstrip('.').upper())
    p=doc.add_paragraph('包内文件：'+relative(ref['source']),'Source Path');p.paragraph_format.keep_with_next=True
    doc.add_paragraph('原始来源：'+ref['source'],'Source Path')

footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 ')
for instr in ['PAGE','NUMPAGES']:
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),instr);footer._p.append(field)
    footer.add_run(' 页 / 共 ' if instr=='PAGE' else ' 页')
for r in footer.runs:r.font.size=Pt(9)
doc.core_properties.title='结构化人—物—场景建模 首轮 S0 S1 验证'
doc.core_properties.subject='BEHAVE 两条开发序列的结构化建模与独立评价'
doc.core_properties.author='HOI 项目'
doc.core_properties.keywords='HOI;3DGS;S0;S1;BEHAVE;2026-09-23'

output=PACKAGE/'REPORT.docx';doc.save(output)

# Embed embeddable TrueType fonts so Chinese and metric symbols travel with file.
NSW='http://schemas.openxmlformats.org/wordprocessingml/2006/main';NSR='http://schemas.openxmlformats.org/officeDocument/2006/relationships';NSPKG='http://schemas.openxmlformats.org/package/2006/relationships'
font_files=[(CJK,'embedRegular',Path('/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf')),(LATIN,'embedRegular',Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')),(LATIN,'embedBold',Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))]
with zipfile.ZipFile(output) as z:parts={n:z.read(n) for n in z.namelist()}
ft=etree.fromstring(parts['word/fontTable.xml']);rels=etree.Element('{'+NSPKG+'}Relationships',nsmap={None:NSPKG});font_records=[]
for k,(family,kind,path) in enumerate(font_files):
    data=path.read_bytes();nt=struct.unpack('>H',data[4:6])[0];tabs={}
    for ix in range(nt):tag,checksum,off,length=struct.unpack('>4sIII',data[12+ix*16:28+ix*16]);tabs[tag.decode('ascii')]=(off,length)
    fs=struct.unpack('>H',data[tabs['OS/2'][0]+8:tabs['OS/2'][0]+10])[0]
    assert not fs&0x2 and not fs&0x200,('Embedding prohibited',path,fs)
    key=uuid.uuid5(uuid.NAMESPACE_URL,str(path));mask=key.bytes[::-1];blob=bytearray(data)
    for ix in range(32):blob[ix]^=mask[ix%16]
    rid='rIdEmbeddedFont'+str(k+1);fname=f'font{k+1}.odttf';parts['word/fonts/'+fname]=bytes(blob)
    f=next((x for x in ft if x.get('{'+NSW+'}name')==family),None)
    if f is None:f=etree.SubElement(ft,'{'+NSW+'}font');f.set('{'+NSW+'}name',family)
    emb=etree.SubElement(f,'{'+NSW+'}'+kind);emb.set('{'+NSR+'}id',rid);emb.set('{'+NSW+'}fontKey','{'+str(key).upper()+'}')
    rel=etree.SubElement(rels,'{'+NSPKG+'}Relationship');rel.set('Id',rid);rel.set('Type',NSR+'/font');rel.set('Target','fonts/'+fname)
    font_records.append({'family':family,'kind':kind,'source':str(path),'fsType':fs,'source_sha256':hashlib.sha256(data).hexdigest(),'embedded_part':'word/fonts/'+fname})
parts['word/fontTable.xml']=etree.tostring(ft,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['word/_rels/fontTable.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=etree.fromstring(parts['[Content_Types].xml']);default=etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default');default.set('Extension','odttf');default.set('ContentType','application/vnd.openxmlformats-officedocument.obfuscatedFont');parts['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
settings=etree.fromstring(parts['word/settings.xml'])
for name in ['embedTrueTypeFonts','saveSubsetFonts']:
    el=etree.SubElement(settings,'{'+NSW+'}'+name);el.set('{'+NSW+'}val','true' if name=='embedTrueTypeFonts' else 'false')
parts['word/settings.xml']=etree.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in parts.items():z.writestr(n,b)

# Verify exact visible source content in source order, permitting added citation IDs.
def norm(s):return re.sub(r'\s+','',s)
def plain(s):
    s=link_re.sub(lambda m:m.group(1),s);s=s.replace('**','').replace('`','');return s
root=etree.fromstring(parts['word/document.xml']);ns={'w':NSW}
alltext=''.join(root.xpath('//w:t/text()',namespaces=ns));alltext=re.sub(r'\s*\[R\d{2}\]','',alltext);cursor=0
for block in block_log:
    values=[v for row in block['cells'] for v in row] if block['kind']=='table' else [block.get('text',block.get('alt',''))]
    for val in values:
        target=norm(plain(val));hay=norm(alltext);at=hay.find(target,cursor)
        assert at>=0,('Missing/reordered source text',target[:160]);cursor=at+len(target)
assert table_count==7
image_parts=[n for n in parts if n.startswith('word/media/')]
for entry in image_log:
    assert any(hashlib.sha256(parts[n]).hexdigest()==entry['sha256'] for n in image_parts),entry
media_external=[]
for name,b in parts.items():
    if name.endswith('.rels'):
        for x in etree.fromstring(b):
            if x.get('Type','').endswith('/image') and x.get('TargetMode')=='External':media_external.append(x.get('Target'))
assert not media_external
rec={'status':'structural_checks_passed_pending_render','source':str(SOURCE),'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'source_characters':len(text),'main_tables':table_count,'main_data_rows':29,'main_headings':9,'source_blocks_verified_in_order':len(block_log),'all_source_text_retained':True,'unique_embedded_images':len(image_parts),'image_placements':image_log,'fonts':font_records,'references':references,'external_image_links':media_external,'output':str(output),'docx_sha256':hashlib.sha256(output.read_bytes()).hexdigest()}
(QA/'conversion_checks.json').write_text(json.dumps(rec,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in rec.items() if k not in ['image_placements','fonts','references']},ensure_ascii=False,indent=2))
